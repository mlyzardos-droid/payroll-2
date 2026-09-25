import sys
import os
import shutil
import sqlite3
import re
from html import escape
from datetime import datetime, timedelta
from openpyxl import Workbook

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

import matplotlib.pyplot as plt
from matplotlib import font_manager as mpl_font_manager

from PySide6.QtCore import QDate, Qt
from PySide6.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QHBoxLayout, QPushButton, QLabel,
    QListWidget, QLineEdit, QMessageBox, QComboBox, QTableWidget,
    QTableWidgetItem, QFileDialog, QDateEdit, QTabWidget,
    QScrollArea, QHeaderView, QSizePolicy, QGridLayout,
    QGroupBox, QToolButton
)


if sys.platform == "win32":
    APP_DATA_ROOT = os.environ.get("LOCALAPPDATA", os.path.expanduser("~"))
else:
    APP_DATA_ROOT = os.path.join(os.path.expanduser("~"), "Library", "Application Support")

APP_DATA_DIR = os.path.join(APP_DATA_ROOT, "PrometheusPayroll")

os.makedirs(APP_DATA_DIR, exist_ok=True)

DB_NAME = os.path.join(APP_DATA_DIR, "payroll.db")
def migrate_legacy_user_files():
    """Copy legacy working-directory data only when the new user database is absent."""
    legacy_db = os.path.join(os.path.dirname(os.path.abspath(__file__)), "payroll.db")
    if not os.path.exists(DB_NAME) and os.path.isfile(legacy_db) and os.path.abspath(legacy_db) != os.path.abspath(DB_NAME):
        shutil.copy2(legacy_db, DB_NAME)
migrate_legacy_user_files()


class PayrollApp(QWidget):
    def __init__(self):
        super().__init__()

        self.editing_record_id = None
        self.editing_period_id = None
        self.last_deleted_employee = None
        self.last_edit_snapshot = None

        self.setWindowTitle("Prometheus Payroll v1.1.2 — Hourly Payroll Edition")
        screen = QApplication.primaryScreen()
        available = screen.availableGeometry() if screen else None
        available_width = available.width() if available else 1440
        available_height = available.height() if available else 900
        minimum_width = max(320, min(720, available_width - 32))
        minimum_height = max(280, min(540, available_height - 48))
        self.setMinimumSize(minimum_width, minimum_height)
        self.resize(
            max(minimum_width, min(1280, available_width - 32)),
            max(minimum_height, min(820, available_height - 48)),
        )

        self.conn = sqlite3.connect(DB_NAME)
        self.cursor = self.conn.cursor()
        self.migration_summary = self.create_tables()

        self.apply_theme()

        main_layout = QVBoxLayout()

        main_layout.addWidget(self.build_tabs_view())

        self.setLayout(main_layout)

        self.load_employees()
        self.load_history()
        self.load_monthly_summary()
        if self.migration_summary:
            QMessageBox.information(self, "Ολοκληρώθηκε η ενημέρωση", self.migration_summary)

    def apply_theme(self):
        self.setStyleSheet("""
            QWidget { background-color: #34383f; color: #f2f3f5; font-size: 14px; }
            QLineEdit, QComboBox, QDateEdit {
                background-color: #454a53; color: #ffffff; border: 1px solid #636a75;
                padding: 6px; border-radius: 6px;
            }
            QPushButton {
                background-color: #3978d4; color: white; padding: 8px;
                border-radius: 8px; font-weight: bold;
            }
            QPushButton:hover { background-color: #4b89e2; }
            QPushButton:disabled { background-color: #575d66; color: #c4c7cc; }
            QTableWidget { background-color: #3b4048; color: #f5f6f7; gridline-color: #5a606a; }
            QTableWidget::item:selected { background-color: #3d72b9; }
            QHeaderView::section {
                background-color: #4a5059; color: #ffffff; padding: 7px; border: 1px solid #626974;
            }
            QListWidget { background-color: #41464e; border: 1px solid #626974; }
            QTabWidget::pane { border: 1px solid #565c65; top: -1px; }
            QTabBar::tab { background: #454a53; color: #f2f3f5; padding: 10px 14px; border-radius: 6px; margin-right: 2px; }
            QTabBar::tab:selected { background: #3978d4; color: white; }
            QGroupBox {
                border: 1px solid #5b616b; border-radius: 8px;
                margin-top: 10px; padding: 10px;
            }
            QGroupBox::title { subcontrol-origin: margin; left: 10px; padding: 0 4px; }
            QLabel#payrollMetricValue { font-size: 19px; font-weight: bold; }
            QScrollArea { border: none; }
        """)

    def build_tabs_view(self):
        tabs = QTabWidget()
        self.tabs = tabs

        def add_scrollable_tab(title):
            page = QWidget()
            page.setMinimumWidth(0)
            area = QScrollArea()
            area.setWidgetResizable(True)
            area.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
            area.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
            area.setWidget(page)
            tabs.addTab(area, title)
            return page

        self.employees_tab = add_scrollable_tab("Employees")
        self.attendance_tab = add_scrollable_tab("Attendance")
        self.payroll_tab = add_scrollable_tab("Payroll")
        self.reports_tab = add_scrollable_tab("Reports")

        self.build_employees_ui(self.employees_tab)
        self.build_attendance_ui(self.attendance_tab)
        self.build_payroll_ui(self.payroll_tab)
        self.build_reports_ui(self.reports_tab)

        return tabs

    @staticmethod
    def configure_responsive_table(table):
        table.setMinimumWidth(0)
        table.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        table.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        table.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        table.setWordWrap(False)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        table.horizontalHeader().setStretchLastSection(True)

    def build_employees_ui(self, parent):
        layout = QVBoxLayout()

        layout.addWidget(QLabel("Διαχείριση Υπαλλήλων"))

        self.employee_input = QLineEdit()
        self.employee_input.setPlaceholderText("Όνομα υπαλλήλου")
        layout.addWidget(self.employee_input)

        self.employee_hourly_rate_input = QLineEdit()
        self.employee_hourly_rate_input.setPlaceholderText("Βασικό ωρομίσθιο (€), π.χ. 7 ή 7,50")
        layout.addWidget(self.employee_hourly_rate_input)

        self.employee_submit_btn = QPushButton("Προσθήκη Υπαλλήλου")
        self.employee_submit_btn.setToolTip("Προσθήκη νέου υπαλλήλου ή αποθήκευση ονόματος και ωρομισθίου του επιλεγμένου υπαλλήλου.")
        self.employee_submit_btn.clicked.connect(self.add_employee)
        layout.addWidget(self.employee_submit_btn)

        new_employee_btn = QPushButton("Νέος Υπάλληλος / Καθαρισμός")
        new_employee_btn.clicked.connect(self.clear_employee_form)
        layout.addWidget(new_employee_btn)

        self.employee_rate_save_btn = QPushButton("Αλλαγή Ωρομισθίου…")
        self.employee_rate_save_btn.setToolTip("Αλλάζει το ωρομίσθιο και σε ρωτά αν θα επανυπολογιστούν οι παλιές ώρες ή αν η αλλαγή θα ισχύσει από σήμερα.")
        self.employee_rate_save_btn.setEnabled(False)
        self.employee_rate_save_btn.clicked.connect(self.save_employee_hourly_rate)
        layout.addWidget(self.employee_rate_save_btn)

        self.employee_list = QListWidget()
        self.employee_list.currentItemChanged.connect(self.load_selected_employee_rate)
        layout.addWidget(self.employee_list)

        remove_btn = QPushButton("Αφαίρεση Υπαλλήλου")
        remove_btn.clicked.connect(self.remove_employee)
        layout.addWidget(remove_btn)

        undo_remove_btn = QPushButton("Αναίρεση Διαγραφής Υπαλλήλου")
        undo_remove_btn.clicked.connect(self.undo_delete_employee)
        layout.addWidget(undo_remove_btn)

        parent.setLayout(layout)

    def build_attendance_ui(self, parent):
        layout = QVBoxLayout()

        layout.addWidget(QLabel("Καταχώρηση Παρουσίας"))

        self.employee_select = QComboBox()
        layout.addWidget(self.employee_select)

        self.date_input = QDateEdit()
        self.date_input.setCalendarPopup(True)
        self.date_input.setDate(QDate.currentDate())
        layout.addWidget(self.date_input)

        self.check_in_input = QLineEdit()
        self.check_in_input.setPlaceholderText("Προσέλευση π.χ. 9 ή 09:00")
        layout.addWidget(self.check_in_input)

        self.check_out_input = QLineEdit()
        self.check_out_input.setPlaceholderText("Αποχώρηση π.χ. 18 ή 18:30")
        layout.addWidget(self.check_out_input)

        self.break_input = QLineEdit()
        self.break_input.setPlaceholderText("Διάλειμμα σε λεπτά")
        layout.addWidget(self.break_input)

        break_note = QLabel(
            "Στην ημερήσια καταχώριση αφαιρείται το διάλειμμα από την προσέλευση–αποχώρηση. "
            "Στη συγκεντρωτική καταχώριση οι καθαρές ώρες και το διάλειμμα δηλώνονται χωριστά."
        )
        break_note.setWordWrap(True)
        layout.addWidget(break_note)

        self.save_btn = QPushButton("Αποθήκευση Ωρών")
        self.save_btn.clicked.connect(self.save_attendance)
        layout.addWidget(self.save_btn)

        self.aggregate_attendance_toggle = QToolButton()
        self.aggregate_attendance_toggle.setText("Καταχώριση συνολικών ωρών για διάστημα")
        self.aggregate_attendance_toggle.setCheckable(True)
        self.aggregate_attendance_toggle.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        self.aggregate_attendance_toggle.setArrowType(Qt.RightArrow)
        self.aggregate_attendance_panel = QWidget()
        aggregate_layout = QGridLayout(self.aggregate_attendance_panel)
        aggregate_layout.setContentsMargins(0, 0, 0, 0)
        aggregate_layout.addWidget(QLabel("Από ημερομηνία"), 0, 0)
        aggregate_layout.addWidget(QLabel("Έως ημερομηνία"), 0, 1)
        self.aggregate_start_date_input = QDateEdit()
        self.aggregate_start_date_input.setCalendarPopup(True)
        self.aggregate_start_date_input.setDate(QDate.currentDate().addDays(-6))
        self.aggregate_end_date_input = QDateEdit()
        self.aggregate_end_date_input.setCalendarPopup(True)
        self.aggregate_end_date_input.setDate(QDate.currentDate())
        aggregate_layout.addWidget(self.aggregate_start_date_input, 1, 0)
        aggregate_layout.addWidget(self.aggregate_end_date_input, 1, 1)
        aggregate_layout.addWidget(QLabel("Καθαρές ώρες εργασίας"), 2, 0)
        aggregate_layout.addWidget(QLabel("Ώρες διαλείμματος (ξεχωριστά)"), 2, 1)
        self.aggregate_work_hours_input = QLineEdit()
        self.aggregate_work_hours_input.setPlaceholderText("π.χ. 30")
        self.aggregate_break_hours_input = QLineEdit()
        self.aggregate_break_hours_input.setPlaceholderText("π.χ. 6")
        aggregate_layout.addWidget(self.aggregate_work_hours_input, 3, 0)
        aggregate_layout.addWidget(self.aggregate_break_hours_input, 3, 1)
        aggregate_note = QLabel(
            "Οι καθαρές ώρες χρησιμοποιούνται στη μισθοδοσία. Οι ώρες διαλείμματος καταγράφονται για αναφορά και δεν αφαιρούνται ξανά. "
            "Αν το διάστημα περνά σε άλλο μήνα, οι καθαρές ώρες επιμερίζονται ανά ημερολογιακή ημέρα."
        )
        aggregate_note.setWordWrap(True)
        aggregate_layout.addWidget(aggregate_note, 4, 0, 1, 2)
        self.aggregate_attendance_save_btn = QPushButton("Αποθήκευση Συγκεντρωτικών Ωρών")
        self.aggregate_attendance_save_btn.clicked.connect(self.save_aggregate_attendance)
        aggregate_layout.addWidget(self.aggregate_attendance_save_btn, 5, 0, 1, 2)
        self.aggregate_attendance_panel.setVisible(False)
        self.aggregate_attendance_toggle.toggled.connect(self.toggle_aggregate_attendance)
        layout.addWidget(self.aggregate_attendance_toggle)
        layout.addWidget(self.aggregate_attendance_panel)

        cancel_btn = QPushButton("Ακύρωση Επεξεργασίας")
        cancel_btn.clicked.connect(self.cancel_edit)
        layout.addWidget(cancel_btn)

        undo_edit_btn = QPushButton("Αναίρεση Τελευταίας Αλλαγής Μισθοδοσίας")
        undo_edit_btn.clicked.connect(self.undo_last_edit)
        layout.addWidget(undo_edit_btn)

        self.result_label = QLabel("")
        layout.addWidget(self.result_label)

        history_box = QGroupBox("Ιστορικό παρουσιών")
        history_layout = QVBoxLayout(history_box)
        history_filters = QGridLayout()
        history_filters.addWidget(QLabel("Υπάλληλος"), 0, 0)
        history_filters.addWidget(QLabel("Από"), 0, 1)
        history_filters.addWidget(QLabel("Έως"), 0, 2)
        self.filter_employee_select = QComboBox()
        self.filter_employee_select.setEditable(True)
        self.filter_employee_select.setInsertPolicy(QComboBox.NoInsert)
        self.filter_employee_select.lineEdit().setPlaceholderText("Όλοι οι υπάλληλοι ή γράψε όνομα")
        self.filter_employee_select.completer().setFilterMode(Qt.MatchContains)
        self.filter_employee_select.completer().setCaseSensitivity(Qt.CaseInsensitive)
        self.filter_employee_select.currentIndexChanged.connect(self.load_history)
        self.filter_employee_select.currentTextChanged.connect(self.load_history)
        history_filters.addWidget(self.filter_employee_select, 1, 0)
        self.from_date_input = QDateEdit()
        self.from_date_input.setCalendarPopup(True)
        self.from_date_input.setDate(self.get_data_start_date())
        self.from_date_input.dateChanged.connect(self.load_history)
        history_filters.addWidget(self.from_date_input, 1, 1)
        self.to_date_input = QDateEdit()
        self.to_date_input.setCalendarPopup(True)
        self.to_date_input.setDate(QDate.currentDate())
        self.to_date_input.dateChanged.connect(self.load_history)
        history_filters.addWidget(self.to_date_input, 1, 2)
        history_filters.setColumnStretch(0, 2)
        history_filters.setColumnStretch(1, 1)
        history_filters.setColumnStretch(2, 1)
        history_layout.addLayout(history_filters)

        self.history_table = QTableWidget()
        self.history_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.history_table.setSelectionMode(QTableWidget.SingleSelection)
        self.history_table.setColumnCount(7)
        self.history_table.setHorizontalHeaderLabels([
            "ID", "Υπάλληλος", "Ημερομηνία", "Προσέλευση", "Αποχώρηση",
            "Διάλειμμα", "Πραγματικές ώρες"
        ])
        self.configure_responsive_table(self.history_table)
        self.history_table.setMinimumHeight(240)
        history_layout.addWidget(self.history_table)
        history_buttons = QHBoxLayout()
        edit_record_btn = QPushButton("Επεξεργασία επιλεγμένης καταχώρισης")
        edit_record_btn.clicked.connect(self.load_selected_record_for_edit)
        delete_record_btn = QPushButton("Διαγραφή επιλεγμένης καταχώρισης")
        delete_record_btn.clicked.connect(self.delete_selected_record)
        clear_history_filters_btn = QPushButton("Καθαρισμός φίλτρων")
        clear_history_filters_btn.clicked.connect(self.clear_filters)
        history_buttons.addWidget(edit_record_btn)
        history_buttons.addWidget(delete_record_btn)
        history_buttons.addWidget(clear_history_filters_btn)
        history_layout.addLayout(history_buttons)
        self.summary_label = QLabel("")
        history_layout.addWidget(self.summary_label)
        layout.addWidget(history_box)

        parent.setLayout(layout)

    def get_data_start_date(self):
        values = []
        for table, column in (("attendance", "date"), ("attendance_periods", "start_date")):
            try:
                values.extend(row[0] for row in self.cursor.execute(
                    f"SELECT {column} FROM {table} WHERE {column} IS NOT NULL"
                ).fetchall())
            except sqlite3.Error:
                pass
        for table in ("monthly_payments", "payroll_payments"):
            try:
                values.extend(
                    f"{int(year):04d}-{int(month):02d}-01"
                    for year, month in self.cursor.execute(
                        f"SELECT year, month FROM {table}"
                    ).fetchall()
                )
            except sqlite3.Error:
                pass
        try:
            values.extend(
                row[0][:10] for row in self.cursor.execute(
                    "SELECT created_at FROM payroll_payments WHERE created_at IS NOT NULL"
                ).fetchall()
            )
        except sqlite3.Error:
            pass
        parsed = []
        for value in values:
            for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
                try:
                    parsed.append(datetime.strptime(str(value), fmt).date())
                    break
                except (TypeError, ValueError):
                    continue
        return QDate(min(parsed).year, min(parsed).month, min(parsed).day) if parsed else QDate.currentDate().addMonths(-1)

    def build_payroll_ui(self, parent):
        layout = QVBoxLayout()
        layout.setSpacing(12)

        summary_box = QGroupBox("Σύνοψη μισθοδοσίας")
        self.employee_summary_group = summary_box
        summary_layout = QVBoxLayout(summary_box)

        self.employee_summary_table = QTableWidget()
        self.employee_summary_table.setColumnCount(7)
        self.employee_summary_table.setHorizontalHeaderLabels([
            "Υπάλληλος", "Πραγματικές Ώρες", "Πληρωμένες Ώρες",
            "Υπόλοιπο Ωρών", "Σύνολο Μισθοδοσίας", "Πληρωμένο Ποσό",
            "Υπόλοιπο (€) · αρνητικό = υπερπληρωμή"
        ])
        self.configure_responsive_table(self.employee_summary_table)
        self.employee_summary_table.setMinimumHeight(180)
        self.employee_summary_table.cellClicked.connect(self.select_monthly_employee_from_row)
        summary_layout.addWidget(self.employee_summary_table)

        monthly_box = QGroupBox("Μηνιαία μισθοδοσία")
        monthly_layout = QVBoxLayout(monthly_box)
        monthly_layout.setSpacing(10)
        selectors = QGridLayout()
        self.payroll_selectors = selectors
        selectors.setHorizontalSpacing(10)
        selectors.addWidget(QLabel("Υπάλληλος"), 0, 0)
        selectors.addWidget(QLabel("Προβολή"), 0, 1)
        self.monthly_month_label = QLabel("Μήνας πληρωμής")
        selectors.addWidget(self.monthly_month_label, 0, 2)
        selectors.addWidget(QLabel("Έτος"), 0, 3)
        self.monthly_employee_select = QComboBox()
        selectors.addWidget(self.monthly_employee_select, 1, 0)
        self.payroll_period_mode = QComboBox()
        self.payroll_period_mode.addItem("Σύνολο", "total")
        self.payroll_period_mode.addItem("Ανά μήνα", "month")
        selectors.addWidget(self.payroll_period_mode, 1, 1)
        self.monthly_month_select = QComboBox()
        for month in range(1, 13):
            self.monthly_month_select.addItem(
                ("Ιανουάριος", "Φεβρουάριος", "Μάρτιος", "Απρίλιος", "Μάιος", "Ιούνιος",
                 "Ιούλιος", "Αύγουστος", "Σεπτέμβριος", "Οκτώβριος", "Νοέμβριος", "Δεκέμβριος")[month - 1],
                month,
            )
        selectors.addWidget(self.monthly_month_select, 1, 2)
        self.monthly_year_input = QLineEdit(str(QDate.currentDate().year()))
        self.monthly_year_input.setPlaceholderText("Έτος")
        self.monthly_year_input.setMaximumWidth(130)
        selectors.addWidget(self.monthly_year_input, 1, 3)
        selectors.setColumnStretch(0, 3)
        selectors.setColumnStretch(1, 2)
        monthly_layout.addLayout(selectors)

        metrics = QHBoxLayout()
        metrics.setSpacing(10)
        self.monthly_total_value = self.create_payroll_metric(metrics, "Σύνολο")
        self.monthly_paid_value = self.create_payroll_metric(metrics, "Πληρωμένο")
        self.monthly_due_value = self.create_payroll_metric(metrics, "Υπόλοιπο · αρνητικό = υπερπληρωμή")
        monthly_layout.addLayout(metrics)

        payment_row = QHBoxLayout()
        self.monthly_paid_hours_input = QLineEdit()
        self.monthly_paid_hours_input.setPlaceholderText("Ώρες που έχουν ήδη πληρωθεί")
        self.monthly_hourly_rate_input = QLineEdit()
        self.monthly_hourly_rate_input.setPlaceholderText("Ωρομίσθιο (€)")
        self.new_payment_amount_input = QLineEdit()
        self.new_payment_amount_input.setPlaceholderText("Πληρωμή τώρα (€), π.χ. 50 ή 50,00")
        self.new_payment_amount_input.setMinimumWidth(160)
        payment_row.addWidget(self.new_payment_amount_input, 1)
        self.add_payment_btn = QPushButton("Καταχώριση Πληρωμής")
        self.add_payment_btn.clicked.connect(self.add_monthly_payment)
        self.add_payment_btn.setMinimumWidth(190)
        self.add_payment_btn.setEnabled(False)
        payment_row.addWidget(self.add_payment_btn)
        monthly_layout.addLayout(payment_row)
        payment_note = QLabel(
            "Η πληρωμή κατανέμεται αυτόματα στους παλαιότερους οφειλόμενους μήνες πρώτα."
        )
        payment_note.setWordWrap(True)
        monthly_layout.addWidget(payment_note)

        self.monthly_adjustments_toggle = QToolButton()
        self.monthly_adjustments_toggle.setText("Καταχώριση ωρών που έχουν ήδη πληρωθεί")
        self.monthly_adjustments_toggle.setCheckable(True)
        self.monthly_adjustments_toggle.setChecked(False)
        self.monthly_adjustments_toggle.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        self.monthly_adjustments_toggle.setArrowType(Qt.RightArrow)
        self.monthly_adjustments_panel = QWidget()
        adjustments_layout = QGridLayout(self.monthly_adjustments_panel)
        adjustments_layout.setContentsMargins(0, 0, 0, 0)
        adjustments_layout.addWidget(QLabel("Ώρες που έχουν ήδη πληρωθεί"), 0, 0)
        adjustments_layout.addWidget(QLabel("Ωρομίσθιο μήνα (€), προαιρετικό"), 0, 1)
        adjustments_layout.addWidget(self.monthly_paid_hours_input, 1, 0)
        adjustments_layout.addWidget(self.monthly_hourly_rate_input, 1, 1)
        paid_hours_note = QLabel(
            "Οι ώρες καλύπτουν πρώτα τις απλήρωτες ώρες των παλαιότερων μηνών. "
            "Όσες περισσέψουν μένουν αδιάθετες για επόμενες ώρες και δεν δημιουργούν αρνητικό υπόλοιπο."
        )
        paid_hours_note.setWordWrap(True)
        adjustments_layout.addWidget(paid_hours_note, 2, 0, 1, 2)
        monthly_save = QPushButton("Αποθήκευση ήδη πληρωμένων ωρών")
        monthly_save.clicked.connect(self.save_monthly_payroll)
        adjustments_layout.addWidget(monthly_save, 3, 0, 1, 2)
        self.monthly_adjustments_panel.setVisible(False)
        self.monthly_adjustments_toggle.toggled.connect(self.toggle_monthly_adjustments)
        monthly_layout.addWidget(self.monthly_adjustments_toggle)
        monthly_layout.addWidget(self.monthly_adjustments_panel)

        self.monthly_month_select.currentIndexChanged.connect(self.refresh_monthly_views)
        self.monthly_year_input.textChanged.connect(self.refresh_monthly_views)
        self.monthly_employee_select.currentIndexChanged.connect(self.refresh_monthly_views)
        self.payroll_period_mode.currentIndexChanged.connect(self.refresh_monthly_views)
        self.monthly_month_select.setVisible(False)
        self.monthly_year_input.setVisible(False)
        selectors.itemAtPosition(0, 2).widget().setVisible(False)
        selectors.itemAtPosition(0, 3).widget().setVisible(False)
        self.monthly_summary_label = QLabel("")
        self.monthly_summary_label.setWordWrap(True)
        monthly_layout.addWidget(self.monthly_summary_label)

        self.payment_history_toggle = QToolButton()
        self.payment_history_toggle.setText("Ιστορικό πληρωμών")
        self.payment_history_toggle.setCheckable(True)
        self.payment_history_toggle.setToolButtonStyle(Qt.ToolButtonTextBesideIcon)
        self.payment_history_toggle.setArrowType(Qt.RightArrow)
        self.payment_history_panel = QWidget()
        payment_history_layout = QVBoxLayout(self.payment_history_panel)
        payment_history_layout.setContentsMargins(0, 0, 0, 0)
        self.payment_history_label = QLabel("")
        payment_history_layout.addWidget(self.payment_history_label)
        self.payment_history_table = QTableWidget()
        self.payment_history_table.setColumnCount(4)
        self.payment_history_table.setHorizontalHeaderLabels([
            "Ημερομηνία καταχώρισης", "Υπάλληλος", "Μήνας που αφορά", "Ποσό"
        ])
        self.configure_responsive_table(self.payment_history_table)
        self.payment_history_table.setMinimumHeight(130)
        self.payment_history_table.setMaximumHeight(230)
        payment_history_layout.addWidget(self.payment_history_table)
        self.payment_history_panel.setVisible(False)
        self.payment_history_toggle.toggled.connect(self.toggle_payment_history)
        monthly_layout.addWidget(self.payment_history_toggle)
        monthly_layout.addWidget(self.payment_history_panel)
        layout.addWidget(monthly_box)
        layout.addWidget(summary_box)

        parent.setLayout(layout)

    @staticmethod
    def create_payroll_metric(layout, title):
        card = QGroupBox()
        card_layout = QVBoxLayout(card)
        card_layout.setContentsMargins(10, 8, 10, 8)
        card_layout.setSpacing(2)
        card_layout.addWidget(QLabel(title))
        value = QLabel("—")
        value.setObjectName("payrollMetricValue")
        card_layout.addWidget(value)
        layout.addWidget(card, 1)
        return value

    def toggle_monthly_adjustments(self, expanded):
        if expanded and self.payroll_period_mode.currentData() == "total":
            month_mode_index = self.payroll_period_mode.findData("month")
            if month_mode_index >= 0:
                self.payroll_period_mode.setCurrentIndex(month_mode_index)
        self.monthly_adjustments_panel.setVisible(expanded)
        self.monthly_adjustments_toggle.setArrowType(Qt.DownArrow if expanded else Qt.RightArrow)

    def toggle_payment_history(self, expanded):
        self.payment_history_panel.setVisible(expanded)
        self.payment_history_toggle.setArrowType(Qt.DownArrow if expanded else Qt.RightArrow)

    def load_payment_history(self):
        if not hasattr(self, "payment_history_table"):
            return
        employee_id = self.monthly_employee_select.currentData()
        if employee_id is None:
            rows = self.cursor.execute("""
                SELECT payments.created_at, COALESCE(employees.name, 'Διαγραμμένος υπάλληλος'),
                       payments.year, payments.month, payments.amount
                FROM payroll_payments AS payments
                LEFT JOIN employees ON employees.id = payments.employee_id
                ORDER BY payments.created_at DESC, payments.id DESC
            """).fetchall()
            title = "Όλοι οι υπάλληλοι"
        else:
            rows = self.cursor.execute("""
                SELECT payments.created_at, COALESCE(employees.name, 'Διαγραμμένος υπάλληλος'),
                       payments.year, payments.month, payments.amount
                FROM payroll_payments AS payments
                LEFT JOIN employees ON employees.id = payments.employee_id
                WHERE payments.employee_id=?
                ORDER BY payments.created_at DESC, payments.id DESC
            """, (employee_id,)).fetchall()
            employee = self.monthly_employee_select.currentText()
            title = employee

        self.payment_history_table.clearContents()
        self.payment_history_table.setRowCount(len(rows))
        total = 0.0
        month_names = ("Ιανουάριος", "Φεβρουάριος", "Μάρτιος", "Απρίλιος", "Μάιος", "Ιούνιος",
                       "Ιούλιος", "Αύγουστος", "Σεπτέμβριος", "Οκτώβριος", "Νοέμβριος", "Δεκέμβριος")
        for row_index, (created_at, name, year, month, amount) in enumerate(rows):
            try:
                recorded_at = datetime.fromisoformat(created_at).strftime("%d/%m/%Y %H:%M")
            except (TypeError, ValueError):
                recorded_at = str(created_at or "")
            month_name = month_names[int(month) - 1] if 1 <= int(month) <= 12 else str(month)
            values = [recorded_at, name, f"{month_name} {year}", f"{float(amount or 0):.2f} €"]
            for column, value in enumerate(values):
                self.payment_history_table.setItem(row_index, column, QTableWidgetItem(str(value)))
            total += float(amount or 0)
        self.payment_history_table.resizeColumnsToContents()
        self.payment_history_label.setText(
            f"{title} · {len(rows)} καταχωρημένες πληρωμές · Σύνολο {total:.2f} €"
            if rows else f"{title} · Δεν έχουν καταχωρηθεί πληρωμές ακόμη."
        )

    def select_monthly_employee_from_row(self, row, _column):
        item = self.employee_summary_table.item(row, 0)
        if item is None:
            return
        employee_id = item.data(Qt.UserRole)
        index = self.monthly_employee_select.findData(employee_id)
        self.monthly_employee_select.blockSignals(True)
        if index >= 0:
            self.monthly_employee_select.setCurrentIndex(index)
        if self.payroll_period_mode.currentData() == "month":
            month_item = self.employee_summary_table.item(row, 1)
            if month_item is not None:
                month_index = self.monthly_month_select.findData(month_item.data(Qt.UserRole))
                if month_index >= 0:
                    self.monthly_month_select.setCurrentIndex(month_index)
        self.monthly_employee_select.blockSignals(False)
        self.refresh_monthly_views()

    def build_reports_ui(self, parent):
        layout = QVBoxLayout()
        layout.addWidget(QLabel("Διάλεξε αναφορά και περίοδο. Οι αναφορές ενημερώνονται αυτόματα με τα φίλτρα."))
        filters = QGridLayout()
        filters.addWidget(QLabel("Αναφορά"), 0, 0)
        filters.addWidget(QLabel("Υπάλληλος"), 0, 1)
        filters.addWidget(QLabel("Από"), 0, 2)
        filters.addWidget(QLabel("Έως"), 0, 3)
        self.report_type_select = QComboBox()
        self.report_type_select.addItem("Σύνοψη μισθοδοσίας", "payroll")
        self.report_type_select.addItem("Παρουσίες αναλυτικά", "attendance")
        self.report_type_select.addItem("Ιστορικό πληρωμών", "payments")
        self.report_type_select.addItem("Οφειλές ανά μήνα", "balances")
        filters.addWidget(self.report_type_select, 1, 0)
        self.report_employee_select = QComboBox()
        filters.addWidget(self.report_employee_select, 1, 1)
        self.report_from_date = QDateEdit()
        self.report_from_date.setCalendarPopup(True)
        self.report_from_date.setDate(self.get_data_start_date())
        filters.addWidget(self.report_from_date, 1, 2)
        self.report_to_date = QDateEdit()
        self.report_to_date.setCalendarPopup(True)
        self.report_to_date.setDate(QDate.currentDate())
        filters.addWidget(self.report_to_date, 1, 3)
        all_time_btn = QPushButton("Όλο το διάστημα")
        all_time_btn.clicked.connect(self.set_report_all_time)
        filters.addWidget(all_time_btn, 2, 0)
        filters.setColumnStretch(0, 2)
        filters.setColumnStretch(1, 2)
        filters.setColumnStretch(2, 1)
        filters.setColumnStretch(3, 1)
        layout.addLayout(filters)

        self.report_period_note = QLabel("")
        self.report_period_note.setWordWrap(True)
        layout.addWidget(self.report_period_note)
        self.report_summary_label = QLabel("")
        self.report_summary_label.setObjectName("payrollPeriodSummary")
        layout.addWidget(self.report_summary_label)

        self.report_table = QTableWidget()
        self.configure_responsive_table(self.report_table)
        self.report_table.setMinimumHeight(360)
        layout.addWidget(self.report_table, 1)

        actions = QHBoxLayout()
        refresh_btn = QPushButton("Ανανέωση αναφοράς")
        refresh_btn.clicked.connect(self.refresh_report)
        excel_btn = QPushButton("Εξαγωγή Excel")
        excel_btn.clicked.connect(self.export_to_excel)
        pdf_btn = QPushButton("Εξαγωγή PDF")
        pdf_btn.clicked.connect(self.export_to_pdf)
        employee_export_btn = QPushButton("Ατομική αναλυτική εξαγωγή Excel")
        employee_export_btn.clicked.connect(self.export_employee_payroll_to_excel)
        backup_btn = QPushButton("Backup βάσης")
        backup_btn.clicked.connect(self.backup_database)
        actions.addWidget(refresh_btn)
        actions.addWidget(excel_btn)
        actions.addWidget(pdf_btn)
        actions.addWidget(employee_export_btn)
        actions.addWidget(backup_btn)
        layout.addLayout(actions)

        self.report_type_select.currentIndexChanged.connect(self.refresh_report)
        self.report_employee_select.currentIndexChanged.connect(self.refresh_report)
        self.report_from_date.dateChanged.connect(self.refresh_report)
        self.report_to_date.dateChanged.connect(self.refresh_report)
        parent.setLayout(layout)

    def set_report_all_time(self):
        self.report_from_date.setDate(self.get_data_start_date())
        self.report_to_date.setDate(QDate.currentDate())
        self.refresh_report()

    def report_months_in_range(self, start_date, end_date):
        start = datetime.strptime(start_date, "%Y-%m-%d").date().replace(day=1)
        end = datetime.strptime(end_date, "%Y-%m-%d").date().replace(day=1)
        months = []
        current = start
        while current <= end:
            months.append((current.year, current.month))
            current = (current.replace(day=28) + timedelta(days=4)).replace(day=1)
        return months

    def refresh_report(self, *_args):
        if not hasattr(self, "report_table"):
            return
        start_date = self.report_from_date.date().toString("yyyy-MM-dd")
        end_date = self.report_to_date.date().toString("yyyy-MM-dd")
        table = self.report_table
        if start_date > end_date:
            table.setRowCount(0)
            self.report_summary_label.setText("Η ημερομηνία έναρξης πρέπει να είναι πριν από την ημερομηνία λήξης.")
            return

        report_type = self.report_type_select.currentData()
        employee_id = self.report_employee_select.currentData()
        employees = self.cursor.execute(
            "SELECT id, name FROM employees WHERE id=? ORDER BY name" if employee_id is not None
            else "SELECT id, name FROM employees ORDER BY name",
            (employee_id,) if employee_id is not None else (),
        ).fetchall()
        headers = []
        rows = []
        note = ""

        if report_type == "attendance":
            headers = ["ID", "Υπάλληλος", "Ημερομηνία", "Προσέλευση", "Αποχώρηση", "Διάλειμμα", "Πραγματικές ώρες"]
            records = self.get_filtered_records(employee_id, start_date, end_date)
            rows = [
                ["Συγκεντρωτική" if int(record[0]) < 0 else str(record[0]), *map(str, record[1:])]
                for record in records
            ]
            hours = sum(float(record[6] or 0) for record in records)
            self.report_summary_label.setText(
                f"{len(records)} καταχωρίσεις · {hours:.2f} ώρες εργασίας"
            )
            note = "Οι ώρες διαλείμματος εμφανίζονται ξεχωριστά και δεν αφαιρούνται ξανά από τις συγκεντρωτικές ώρες."

        elif report_type in ("payroll", "balances"):
            months = self.report_months_in_range(start_date, end_date)
            note = "Η μισθοδοσία υπολογίζεται ανά πλήρη ημερολογιακό μήνα που τέμνει το επιλεγμένο διάστημα."
            if report_type == "payroll":
                headers = ["Υπάλληλος", "Πραγματικές ώρες", "Σύνολο μισθοδοσίας", "Ώρες ήδη πληρωμένες", "Συνολικά πληρωμένο", "Υπόλοιπο"]
                total_hours = total_gross = total_paid_hours = total_paid_amount = total_due = 0.0
                for current_id, name in employees:
                    accum = {key: 0.0 for key in ("actual_hours", "gross_pay", "paid_hours", "total_paid_amount", "amount_due")}
                    for year, month in months:
                        payroll = self.calculate_monthly_payroll(current_id, year, month)
                        for key in accum:
                            accum[key] += payroll[key]
                    if not any(accum.values()):
                        continue
                    rows.append([name, f"{accum['actual_hours']:.2f}", f"{accum['gross_pay']:.2f} €",
                                 f"{accum['paid_hours']:.2f}", f"{accum['total_paid_amount']:.2f} €",
                                 f"{accum['amount_due']:.2f} €"])
                    total_hours += accum["actual_hours"]
                    total_gross += accum["gross_pay"]
                    total_paid_hours += accum["paid_hours"]
                    total_paid_amount += accum["total_paid_amount"]
                    total_due += accum["amount_due"]
                self.report_summary_label.setText(
                    f"{len(rows)} υπάλληλοι · {total_hours:.2f} ώρες · Μισθοδοσία {total_gross:.2f} € · "
                    f"Πληρωμένα {total_paid_amount:.2f} € · Υπόλοιπο {total_due:.2f} €"
                )
            else:
                headers = ["Υπάλληλος", "Μήνας", "Ώρες", "Ώρες ήδη πληρωμένες", "Μισθοδοσία", "Πληρωμές που καταχωρίστηκαν", "Πληρωμένο σύνολο", "Κατάσταση", "Υπόλοιπο"]
                month_names = ("Ιανουάριος", "Φεβρουάριος", "Μάρτιος", "Απρίλιος", "Μάιος", "Ιούνιος",
                               "Ιούλιος", "Αύγουστος", "Σεπτέμβριος", "Οκτώβριος", "Νοέμβριος", "Δεκέμβριος")
                total_due = 0.0
                for current_id, name in employees:
                    for year, month in months:
                        payroll = self.calculate_monthly_payroll(current_id, year, month)
                        if not any((payroll["actual_hours"], payroll["paid_hours"], payroll["total_paid_amount"], payroll["amount_due"])):
                            continue
                        payments = self.cursor.execute("""
                            SELECT created_at, amount FROM payroll_payments
                            WHERE employee_id=? AND year=? AND month=? ORDER BY created_at, id
                        """, (current_id, year, month)).fetchall()
                        details = []
                        for created_at, amount in payments:
                            try:
                                date_label = datetime.fromisoformat(created_at).strftime("%d/%m/%Y")
                            except (TypeError, ValueError):
                                date_label = str(created_at or "")
                            details.append(f"{date_label}: {float(amount or 0):.2f} €")
                        hours_amount = payroll["paid_hours"] * payroll["hourly_rate"]
                        if hours_amount:
                            details.insert(0, f"Πληρωμένες ώρες: {payroll['paid_hours']:.2f} · αξία {hours_amount:.2f} €")
                        if payroll["amount_due"] < -0.005:
                            status = "Υπερπληρωμή"
                        elif payroll["amount_due"] <= 0.005:
                            status = "Καλύφθηκε"
                        elif payroll["total_paid_amount"] > 0:
                            status = "Μερική κάλυψη"
                        else:
                            status = "Ακάλυπτος"
                        rows.append([name, f"{month_names[month - 1]} {year}", f"{payroll['actual_hours']:.2f}",
                                     f"{payroll['paid_hours']:.2f}", f"{payroll['gross_pay']:.2f} €", "\n".join(details) or "Καμία πληρωμή",
                                     f"{payroll['total_paid_amount']:.2f} €", status,
                                     f"{payroll['amount_due']:.2f} €"])
                        total_due += payroll["amount_due"]
                self.report_summary_label.setText(
                    f"{len(rows)} μηνιαίες εγγραφές · Συνολικό υπόλοιπο {total_due:.2f} €"
                )

        else:  # payment history
            headers = ["Ημερομηνία καταχώρισης", "Υπάλληλος", "Μήνας", "Είδος", "Ποσό / ώρες", "Αξία (€)"]
            note = "Οι πληρωμές ποσού εμφανίζονται στον μήνα στον οποίο κατανεμήθηκαν. Οι πληρωμένες ώρες είναι οι τελευταίες αποθηκευμένες τιμές ανά μήνα."
            cash_query = """
                SELECT payments.created_at, employees.name, payments.year, payments.month, payments.amount
                FROM payroll_payments AS payments
                JOIN employees ON employees.id=payments.employee_id
                WHERE date(payments.created_at)>=? AND date(payments.created_at)<=?
            """
            params = [start_date, end_date]
            if employee_id is not None:
                cash_query += " AND payments.employee_id=?"
                params.append(employee_id)
            cash_query += " ORDER BY payments.created_at DESC, payments.id DESC"
            payment_rows = self.cursor.execute(cash_query, params).fetchall()
            hours_query = """
                SELECT monthly.updated_at, employees.name, monthly.year, monthly.month,
                       monthly.paid_hours, monthly.hourly_rate
                FROM monthly_payments AS monthly
                JOIN employees ON employees.id=monthly.employee_id
                WHERE monthly.paid_hours>0
                  AND date(monthly.updated_at)>=? AND date(monthly.updated_at)<=?
            """
            hours_params = [start_date, end_date]
            if employee_id is not None:
                hours_query += " AND monthly.employee_id=?"
                hours_params.append(employee_id)
            hours_query += " ORDER BY monthly.updated_at DESC, monthly.id DESC"
            paid_hour_rows = self.cursor.execute(hours_query, hours_params).fetchall()
            month_names = ("Ιανουάριος", "Φεβρουάριος", "Μάρτιος", "Απρίλιος", "Μάιος", "Ιούνιος",
                           "Ιούλιος", "Αύγουστος", "Σεπτέμβριος", "Οκτώβριος", "Νοέμβριος", "Δεκέμβριος")
            for created_at, name, year, month, amount in payment_rows:
                try:
                    date_label = datetime.fromisoformat(created_at).strftime("%d/%m/%Y %H:%M")
                except (TypeError, ValueError):
                    date_label = str(created_at or "")
                rows.append([date_label, name, f"{month_names[int(month) - 1]} {year}",
                             "Πληρωμή ποσού", f"{float(amount or 0):.2f} €", f"{float(amount or 0):.2f} €"])
            for updated_at, name, year, month, paid_hours, rate in paid_hour_rows:
                try:
                    date_label = datetime.fromisoformat(updated_at).strftime("%d/%m/%Y %H:%M")
                except (TypeError, ValueError):
                    date_label = str(updated_at or "")
                value = float(paid_hours or 0) * float(rate or 0)
                rows.append([date_label, name, f"{month_names[int(month) - 1]} {year}",
                             "Ώρες δηλωμένες πληρωμένες", f"{float(paid_hours or 0):.2f} ώρες", f"{value:.2f} €"])
            def payment_date_key(row):
                for fmt in ("%d/%m/%Y %H:%M", "%d/%m/%Y"):
                    try:
                        return datetime.strptime(row[0], fmt)
                    except (TypeError, ValueError):
                        pass
                return datetime.min
            rows.sort(key=payment_date_key, reverse=True)
            total_amount = sum(float(row[4] or 0) for row in payment_rows) + sum(
                float(row[4] or 0) * float(row[5] or 0) for row in paid_hour_rows
            )
            self.report_summary_label.setText(
                f"{len(rows)} καταχωρήσεις · Πληρωμένα ποσά και ώρες συνολικής αξίας {total_amount:.2f} €"
            )

        self.report_period_note.setText(note)
        self.current_report_headers = headers
        self.current_report_rows = rows
        table.clearContents()
        table.setColumnCount(len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.setRowCount(len(rows))
        for row_index, values in enumerate(rows):
            for column, value in enumerate(values):
                table.setItem(row_index, column, QTableWidgetItem(str(value)))
            if report_type == "balances" and len(rows[row_index]) > 4:
                detail_lines = str(rows[row_index][4]).count("\n") + 1
                table.setRowHeight(row_index, max(32, detail_lines * 20))
        table.resizeColumnsToContents()
        table.horizontalHeader().setStretchLastSection(True)
        if not rows:
            self.report_summary_label.setText("Δεν βρέθηκαν δεδομένα για τα φίλτρα που επέλεξες.")

    def create_tables(self):
        self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS employees (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL
            )
        """)

        self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                employee_id INTEGER,
                date TEXT,
                check_in TEXT,
                check_out TEXT,
                break_minutes INTEGER,
                total_hours REAL,
                paid_hours REAL,
                unpaid_hours REAL,
                overtime_hours REAL,
                unpaid_overtime_rate REAL DEFAULT 0,
                unpaid_overtime_amount REAL DEFAULT 0,
                unpaid_overtime_paid INTEGER DEFAULT 0,
                overtime_payment_date TEXT
            )
        """)
        self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS attendance_periods (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                employee_id INTEGER NOT NULL,
                start_date TEXT NOT NULL,
                end_date TEXT NOT NULL,
                total_work_hours REAL NOT NULL CHECK(total_work_hours >= 0),
                break_hours REAL NOT NULL CHECK(break_hours >= 0),
                created_at TEXT NOT NULL,
                CHECK(end_date >= start_date),
                FOREIGN KEY(employee_id) REFERENCES employees(id)
            )
        """)
        self.cursor.execute("""
            CREATE TABLE IF NOT EXISTS employee_hourly_rates (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                employee_id INTEGER NOT NULL,
                effective_from TEXT NOT NULL,
                hourly_rate REAL NOT NULL CHECK(hourly_rate >= 0),
                created_at TEXT NOT NULL,
                UNIQUE(employee_id, effective_from),
                FOREIGN KEY(employee_id) REFERENCES employees(id)
            )
        """)
        self.conn.commit()
        columns = self.table_columns("attendance")
        required = {"unpaid_overtime_amount", "unpaid_overtime_paid", "overtime_payment_date", "overtime_hours", "unpaid_overtime_rate"}
        employee_columns = self.table_columns("employees")
        employee_rate_missing = "hourly_rate" not in employee_columns
        monthly_exists = self.cursor.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='monthly_payments'"
        ).fetchone() is not None
        payments_exists = self.cursor.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='payroll_payments'"
        ).fetchone() is not None
        current_version = self.get_schema_version()
        if required.issubset(columns) and monthly_exists and payments_exists and not employee_rate_missing and current_version >= 4:
            self.cursor.execute("""
                INSERT OR IGNORE INTO employee_hourly_rates(employee_id, effective_from, hourly_rate, created_at)
                SELECT id, '0001-01-01', hourly_rate, ? FROM employees
            """, (datetime.now().isoformat(timespec="seconds"),))
            self.conn.commit()
            return ""

        backup_path = self.create_pre_migration_backup()
        changed_rows = malformed_rows = 0
        recalculate_history = current_version < 2
        self.cursor.execute("BEGIN IMMEDIATE")
        try:
            if employee_rate_missing:
                self.cursor.execute("ALTER TABLE employees ADD COLUMN hourly_rate REAL NOT NULL DEFAULT 0")
            for name, definition in [
                ("unpaid_overtime_amount", "REAL DEFAULT 0"),
                ("unpaid_overtime_paid", "INTEGER DEFAULT 0"),
                ("overtime_payment_date", "TEXT"),
                ("overtime_hours", "REAL DEFAULT 0"),
                ("unpaid_overtime_rate", "REAL DEFAULT 0"),
            ]:
                if name not in self.table_columns("attendance"):
                    self.cursor.execute(f"ALTER TABLE attendance ADD COLUMN {name} {definition}")
            self.cursor.execute("""
                CREATE TABLE IF NOT EXISTS monthly_payments (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    employee_id INTEGER NOT NULL,
                    year INTEGER NOT NULL,
                    month INTEGER NOT NULL,
                    paid_hours REAL NOT NULL DEFAULT 0,
                    hourly_rate REAL NOT NULL DEFAULT 0,
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    UNIQUE(employee_id, year, month),
                    FOREIGN KEY(employee_id) REFERENCES employees(id)
                )
            """)
            self.cursor.execute("""
                CREATE TABLE IF NOT EXISTS payroll_payments (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    employee_id INTEGER NOT NULL,
                    year INTEGER NOT NULL,
                    month INTEGER NOT NULL,
                    amount REAL NOT NULL CHECK(amount > 0),
                    created_at TEXT NOT NULL,
                    FOREIGN KEY(employee_id) REFERENCES employees(id)
                )
            """)
            if recalculate_history:
                rows = self.cursor.execute("SELECT id, date, check_in, check_out, break_minutes FROM attendance").fetchall()
                for row_id, date_value, check_in, check_out, break_minutes in rows:
                    try:
                        start = datetime.strptime(self.normalize_time(check_in or ""), "%H:%M")
                        end = datetime.strptime(self.normalize_time(check_out or ""), "%H:%M")
                        if end <= start:
                            from datetime import timedelta
                            end += timedelta(days=1)
                        actual = (end - start).total_seconds() / 3600 - int(break_minutes or 0) / 60
                        parsed_date = datetime.strptime(str(date_value), "%Y-%m-%d") if "-" in str(date_value) else datetime.strptime(str(date_value), "%Y/%m/%d")
                        if actual < 0:
                            raise ValueError
                    except (ValueError, TypeError):
                        malformed_rows += 1
                        continue
                    self.cursor.execute(
                        "UPDATE attendance SET date=?, total_hours=?, unpaid_hours=MAX(0, ? - COALESCE(paid_hours, 0)), overtime_hours=0 WHERE id=?",
                        (parsed_date.strftime("%Y-%m-%d"), actual, actual, row_id),
                    )
                    changed_rows += 1
            self.cursor.execute("PRAGMA user_version = 4")
            self.conn.commit()
        except Exception:
            self.conn.rollback()
            raise
        self.cursor.execute("""
            INSERT OR IGNORE INTO employee_hourly_rates(employee_id, effective_from, hourly_rate, created_at)
            SELECT id, '0001-01-01', hourly_rate, ? FROM employees
        """, (datetime.now().isoformat(timespec="seconds"),))
        self.conn.commit()
        migration_details = (
            f"Επαναϋπολογίστηκαν {changed_rows} παρουσίες και παραλείφθηκαν {malformed_rows} μη έγκυρες εγγραφές."
            if recalculate_history else "Το ιστορικό παρουσιών διατηρήθηκε ανέπαφο."
        )
        return f"Η βάση ενημερώθηκε. Αντίγραφο ασφαλείας: {backup_path}\n{migration_details}"

    def table_columns(self, table_name):
        self.cursor.execute(f"PRAGMA table_info({table_name})")
        return {column[1] for column in self.cursor.fetchall()}

    def get_schema_version(self):
        return self.cursor.execute("PRAGMA user_version").fetchone()[0]

    def create_pre_migration_backup(self):
        os.makedirs(APP_DATA_DIR, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = os.path.join(APP_DATA_DIR, f"payroll_backup_before_v1_1_0_{timestamp}.db")
        self.conn.commit()
        destination = sqlite3.connect(backup_path)
        try:
            self.conn.backup(destination)
        finally:
            destination.close()
        return backup_path

    def add_column_if_missing(self, table_name, column_name, column_type):
        self.cursor.execute(f"PRAGMA table_info({table_name})")
        columns = [column[1] for column in self.cursor.fetchall()]

        if column_name not in columns:
            self.cursor.execute(f"ALTER TABLE {table_name} ADD COLUMN {column_name} {column_type}")

    def normalize_time(self, value):
        value = value.strip()
        match = re.fullmatch(r"(\d{1,2})(?::(\d{1,2}))?", value)
        if not match:
            raise ValueError
        hour, minute = int(match.group(1)), int(match.group(2) or 0)
        if hour > 23 or minute > 59:
            raise ValueError
        return f"{hour:02d}:{minute:02d}"

    def parse_number(self, value):
        value = value.strip().replace(",", ".")
        if value == "":
            return 0
        return float(value)

    def load_employees(self, selected_employee_id=None):
        current_item = self.employee_list.currentItem() if self.employee_list.count() else None
        current_employee_id = current_item.data(32) if current_item else None
        selected_employee_id = selected_employee_id or current_employee_id
        current_filter_id = self.filter_employee_select.currentData() if self.filter_employee_select.count() > 0 else None
        current_monthly_id = self.monthly_employee_select.currentData() if hasattr(self, "monthly_employee_select") else None
        current_report_id = self.report_employee_select.currentData() if hasattr(self, "report_employee_select") else None

        self.employee_list.clear()
        self.employee_select.clear()
        if hasattr(self, "monthly_employee_select"):
            self.monthly_employee_select.clear()
            self.monthly_employee_select.addItem("Όλοι οι υπάλληλοι", None)

        self.filter_employee_select.blockSignals(True)
        self.filter_employee_select.clear()
        self.filter_employee_select.addItem("Όλοι οι υπάλληλοι", None)
        if hasattr(self, "report_employee_select"):
            self.report_employee_select.blockSignals(True)
            self.report_employee_select.clear()
            self.report_employee_select.addItem("Όλοι οι υπάλληλοι", None)

        self.cursor.execute("SELECT id, name, hourly_rate FROM employees ORDER BY name")
        employees = self.cursor.fetchall()

        selected_index = 0
        monthly_selected_index = 0
        report_selected_index = 0

        for employee_id, name, hourly_rate in employees:
            self.employee_list.addItem(f"{name} — {float(hourly_rate or 0):.2f} €/ώρα")
            self.employee_list.item(self.employee_list.count() - 1).setData(32, employee_id)
            self.employee_list.item(self.employee_list.count() - 1).setToolTip(f"Ωρομίσθιο: {float(hourly_rate or 0):.2f} €")
            self.employee_select.addItem(name, employee_id)
            self.filter_employee_select.addItem(name, employee_id)
            if hasattr(self, "report_employee_select"):
                self.report_employee_select.addItem(name, employee_id)
                if employee_id == current_report_id:
                    report_selected_index = self.report_employee_select.count() - 1
            if hasattr(self, "monthly_employee_select"):
                self.monthly_employee_select.addItem(name, employee_id)
                if employee_id == current_monthly_id:
                    monthly_selected_index = self.monthly_employee_select.count() - 1
            if current_filter_id == employee_id:
                selected_index = self.filter_employee_select.count() - 1

        if self.employee_list.count() > 0:
            selected_row = next(
                (row for row in range(self.employee_list.count())
                 if self.employee_list.item(row).data(32) == selected_employee_id),
                0,
            )
            self.employee_list.setCurrentRow(selected_row)

        self.filter_employee_select.setCurrentIndex(selected_index)
        self.filter_employee_select.blockSignals(False)
        if hasattr(self, "report_employee_select"):
            self.report_employee_select.blockSignals(True)
            self.report_employee_select.setCurrentIndex(report_selected_index)
            self.report_employee_select.blockSignals(False)
        if hasattr(self, "monthly_employee_select"):
            self.monthly_employee_select.blockSignals(True)
            self.monthly_employee_select.setCurrentIndex(monthly_selected_index)
            self.monthly_employee_select.blockSignals(False)
        self.load_monthly_summary()
        self.load_history()

    def load_selected_employee_rate(self, item, _previous=None):
        if item is None:
            self.editing_employee_id = None
            self.employee_input.clear()
            self.employee_hourly_rate_input.clear()
            if hasattr(self, "employee_submit_btn"):
                self.employee_submit_btn.setText("Προσθήκη Υπαλλήλου")
            if hasattr(self, "employee_rate_save_btn"):
                self.employee_rate_save_btn.setEnabled(False)
            return
        row = self.cursor.execute(
            "SELECT name, hourly_rate FROM employees WHERE id=?", (item.data(32),)
        ).fetchone()
        if row:
            self.editing_employee_id = item.data(32)
            self.employee_input.setText(row[0])
            self.employee_hourly_rate_input.setText(str(float(row[1] or 0)))
            self.employee_submit_btn.setText("Αποθήκευση Ονόματος και Ωρομισθίου")
            self.employee_rate_save_btn.setEnabled(True)

    def clear_employee_form(self):
        self.employee_list.setCurrentRow(-1)
        self.editing_employee_id = None
        self.employee_input.clear()
        self.employee_hourly_rate_input.clear()
        self.employee_submit_btn.setText("Προσθήκη Υπαλλήλου")
        self.employee_rate_save_btn.setEnabled(False)

    def save_employee_hourly_rate(self):
        item = self.employee_list.currentItem()
        if item is None:
            QMessageBox.warning(self, "Σφάλμα", "Επίλεξε υπάλληλο για να αποθηκεύσεις το ωρομίσθιο.")
            return
        try:
            hourly_rate = self.parse_number(self.employee_hourly_rate_input.text())
            if hourly_rate < 0:
                raise ValueError
        except ValueError:
            QMessageBox.warning(self, "Σφάλμα", "Μη έγκυρο ωρομίσθιο. Χρησιμοποίησε αριθμό όπως 7, 7.5 ή 7,5.")
            return
        employee_id = item.data(32)
        employee_name = item.text().split(" — ", 1)[0]
        if not self.apply_employee_rate_change(employee_id, hourly_rate):
            return
        self.conn.commit()
        self.load_employees()
        self.load_monthly_summary()
        self.load_history()
        QMessageBox.information(self, "Αποθηκεύτηκε", f"Ωρομίσθιο {hourly_rate:.2f} € αποθηκεύτηκε για τον/την {employee_name}.")

    def apply_employee_rate_change(self, employee_id, hourly_rate):
        row = self.cursor.execute(
            "SELECT hourly_rate FROM employees WHERE id=?", (employee_id,)
        ).fetchone()
        old_rate = float(row[0] or 0) if row else 0.0
        if abs(old_rate - hourly_rate) < 0.000001:
            return True
        dialog = QMessageBox(self)
        dialog.setWindowTitle("Πώς να εφαρμοστεί το νέο ωρομίσθιο;")
        dialog.setIcon(QMessageBox.Question)
        dialog.setText(
            "Επίλεξε πώς θα αλλάξει ο υπολογισμός μισθοδοσίας για τον/την υπάλληλο."
        )
        recalculate_button = dialog.addButton(
            "Επανυπολογισμός όλων των ωρών", QMessageBox.YesRole
        )
        future_button = dialog.addButton(
            "Ισχύει από σήμερα και μετά", QMessageBox.NoRole
        )
        cancel_button = dialog.addButton("Ακύρωση", QMessageBox.RejectRole)
        dialog.setDefaultButton(future_button)
        dialog.exec()
        choice = dialog.clickedButton()
        if choice == cancel_button:
            return False
        now = datetime.now().isoformat(timespec="seconds")
        if choice == recalculate_button:
            self.cursor.execute(
                "UPDATE employees SET hourly_rate=? WHERE id=?", (hourly_rate, employee_id)
            )
            self.cursor.execute(
                "UPDATE employee_hourly_rates SET hourly_rate=? WHERE employee_id=?",
                (hourly_rate, employee_id),
            )
            self.cursor.execute(
                "INSERT OR IGNORE INTO employee_hourly_rates(employee_id, effective_from, hourly_rate, created_at) "
                "VALUES (?, '0001-01-01', ?, ?)", (employee_id, hourly_rate, now)
            )
            self.cursor.execute(
                "UPDATE monthly_payments SET hourly_rate=?, updated_at=? WHERE employee_id=?",
                (hourly_rate, now, employee_id),
            )
        elif choice == future_button:
            self.cursor.execute(
                "UPDATE employees SET hourly_rate=? WHERE id=?", (hourly_rate, employee_id)
            )
            today = QDate.currentDate().toString("yyyy-MM-dd")
            self.cursor.execute("""
                INSERT INTO employee_hourly_rates(employee_id, effective_from, hourly_rate, created_at)
                VALUES (?, ?, ?, ?)
                ON CONFLICT(employee_id, effective_from) DO UPDATE SET
                    hourly_rate=excluded.hourly_rate, created_at=excluded.created_at
            """, (employee_id, today, hourly_rate, now))
        else:
            return False
        return True

    def hourly_rate_for_date(self, employee_id, date_value):
        row = self.cursor.execute("""
            SELECT hourly_rate FROM employee_hourly_rates
            WHERE employee_id=? AND effective_from<=?
            ORDER BY effective_from DESC LIMIT 1
        """, (employee_id, date_value)).fetchone()
        if row:
            return float(row[0] or 0)
        row = self.cursor.execute(
            "SELECT hourly_rate FROM employees WHERE id=?", (employee_id,)
        ).fetchone()
        return float(row[0] or 0) if row else 0.0

    def add_employee(self):
        name = self.employee_input.text().strip()

        if name == "":
            QMessageBox.warning(self, "Σφάλμα", "Δώσε όνομα υπαλλήλου")
            return

        try:
            hourly_rate = self.parse_number(self.employee_hourly_rate_input.text())
            if hourly_rate < 0:
                raise ValueError
        except ValueError:
            QMessageBox.warning(self, "Σφάλμα", "Μη έγκυρο ωρομίσθιο. Χρησιμοποίησε αριθμό όπως 7, 7.5 ή 7,5.")
            return

        if self.editing_employee_id is not None:
            employee_id = self.editing_employee_id
            if not self.apply_employee_rate_change(employee_id, hourly_rate):
                return
            self.cursor.execute("UPDATE employees SET name=? WHERE id=?", (name, employee_id))
            saved_message = f"Τα στοιχεία του/της {name} αποθηκεύτηκαν."
        else:
            self.cursor.execute("INSERT INTO employees (name, hourly_rate) VALUES (?, ?)", (name, hourly_rate))
            employee_id = self.cursor.lastrowid
            self.cursor.execute(
                "INSERT INTO employee_hourly_rates(employee_id, effective_from, hourly_rate, created_at) VALUES (?, '0001-01-01', ?, ?)",
                (employee_id, hourly_rate, datetime.now().isoformat(timespec="seconds")),
            )
            saved_message = f"Ο/Η {name} προστέθηκε."
        self.conn.commit()
        self.load_employees(selected_employee_id=employee_id)
        self.load_history()
        QMessageBox.information(self, "Αποθηκεύτηκε", saved_message)

    def remove_employee(self):
        selected_item = self.employee_list.currentItem()

        if selected_item is None:
            QMessageBox.warning(self, "Σφάλμα", "Επίλεξε υπάλληλο")
            return

        employee_id = selected_item.data(32)
        self.cursor.execute("SELECT id, name, hourly_rate FROM employees WHERE id = ?", (employee_id,))
        employee = self.cursor.fetchone()

        if employee is None:
            QMessageBox.warning(self, "Σφάλμα", "Δεν βρέθηκε ο υπάλληλος")
            return

        employee_id, employee_name, hourly_rate = employee

        self.cursor.execute("SELECT COUNT(*) FROM attendance WHERE employee_id = ?", (employee_id,))
        attendance_count = self.cursor.fetchone()[0]
        period_count = self.cursor.execute(
            "SELECT COUNT(*) FROM attendance_periods WHERE employee_id = ?", (employee_id,)
        ).fetchone()[0]

        if attendance_count > 0 or period_count > 0:
            QMessageBox.warning(
                self,
                "Δεν επιτρέπεται",
                "Ο υπάλληλος έχει καταχωρήσεις μισθοδοσίας. Δεν διαγράφεται για να μη χαθεί ιστορικό."
            )
            return

        confirm = QMessageBox.question(self, "Επιβεβαίωση", f"Να αφαιρεθεί ο υπάλληλος {employee_name};")

        if confirm != QMessageBox.Yes:
            return

        self.last_deleted_employee = {
            "id": employee_id,
            "name": employee_name,
            "hourly_rate": hourly_rate,
        }

        self.cursor.execute("DELETE FROM employees WHERE id = ?", (employee_id,))
        self.conn.commit()

        self.load_employees()
        self.load_history()

        QMessageBox.information(
            self,
            "Διαγράφηκε",
            "Ο υπάλληλος διαγράφηκε. Μπορείς να πατήσεις Αναίρεση Διαγραφής Υπαλλήλου."
        )

    def undo_delete_employee(self):
        if not self.last_deleted_employee:
            QMessageBox.information(self, "Αναίρεση", "Δεν υπάρχει πρόσφατη διαγραφή υπαλλήλου για αναίρεση.")
            return

        employee_id = self.last_deleted_employee["id"]
        employee_name = self.last_deleted_employee["name"]

        self.cursor.execute("SELECT id FROM employees WHERE id = ?", (employee_id,))
        existing = self.cursor.fetchone()

        if existing:
            QMessageBox.warning(self, "Σφάλμα", "Ο υπάλληλος υπάρχει ήδη στη βάση.")
            self.last_deleted_employee = None
            return

        self.cursor.execute(
            "INSERT INTO employees (id, name, hourly_rate) VALUES (?, ?, ?)",
            (employee_id, employee_name, self.last_deleted_employee.get("hourly_rate", 0))
        )
        self.conn.commit()

        self.last_deleted_employee = None
        self.load_employees()
        self.load_history()

        QMessageBox.information(self, "Αναίρεση", f"Ο υπάλληλος {employee_name} επανήλθε.")

    def calculate_attendance_values(self):
        check_in_text = self.check_in_input.text().strip()
        check_out_text = self.check_out_input.text().strip()
        break_text = self.break_input.text().strip()

        if check_in_text == "" or check_out_text == "":
            raise ValueError("missing_time")

        try:
            check_in = self.normalize_time(check_in_text)
            check_out = self.normalize_time(check_out_text)
            start = datetime.strptime(check_in, "%H:%M")
            end = datetime.strptime(check_out, "%H:%M")
        except Exception:
            raise ValueError("invalid_time")

        try:
            break_minutes = int(break_text) if break_text else 0
        except Exception:
            raise ValueError("invalid_number")

        if break_minutes < 0:
            raise ValueError("negative_number")

        paid_hours = 0

        unpaid_overtime_paid = 0
        overtime_payment_date = None

        if end <= start:
            from datetime import timedelta
            end += timedelta(days=1)
        actual_hours = ((end - start).total_seconds() / 60 - break_minutes) / 60

        if actual_hours < 0:
            raise ValueError("negative_hours")

        total_hours = actual_hours
        overtime_hours = 0
        unpaid_hours = max(0, total_hours - paid_hours)
        unpaid_overtime_amount = 0

        return {
            "check_in": check_in,
            "check_out": check_out,
            "break_minutes": break_minutes,
            "paid_hours": paid_hours,
            "unpaid_hours": unpaid_hours,
            "total_hours": total_hours,
            "overtime_hours": overtime_hours,
            "unpaid_overtime_amount": unpaid_overtime_amount,
            "unpaid_overtime_paid": unpaid_overtime_paid,
            "overtime_payment_date": overtime_payment_date
        }

    def save_attendance(self):
        employee_id = self.employee_select.currentData()
        date = self.date_input.date().toString("yyyy-MM-dd")

        if employee_id is None:
            QMessageBox.warning(self, "Σφάλμα", "Δεν υπάρχει υπάλληλος")
            return

        try:
            values = self.calculate_attendance_values()
        except ValueError as error:
            error_text = str(error)

            if error_text == "missing_time":
                QMessageBox.warning(self, "Σφάλμα", "Συμπλήρωσε προσέλευση και αποχώρηση.")
            elif error_text == "invalid_time":
                QMessageBox.warning(self, "Σφάλμα", "Μη έγκυρη ώρα. Βάλε ώρα όπως 9, 09:00 ή 9:30.")
            elif error_text == "invalid_number":
                QMessageBox.warning(self, "Σφάλμα", "Μη έγκυρος αριθμός. Βάλε τιμή όπως 8, 8.5 ή 8,5.")
            elif error_text == "negative_number":
                QMessageBox.warning(self, "Σφάλμα", "Οι αριθμοί δεν μπορούν να είναι αρνητικοί.")
            elif error_text == "negative_hours":
                QMessageBox.warning(self, "Σφάλμα", "Η αποχώρηση πρέπει να είναι μετά την προσέλευση.")
            else:
                QMessageBox.warning(self, "Σφάλμα", "Έλεγξε ώρες και αριθμούς.")
            return

        if self.editing_record_id is None:
            self.cursor.execute("""
                INSERT INTO attendance (
                    employee_id, date, check_in, check_out, break_minutes,
                    total_hours, paid_hours, unpaid_hours, overtime_hours,
                    unpaid_overtime_rate, unpaid_overtime_amount, unpaid_overtime_paid,
                    overtime_payment_date
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                employee_id, date, values["check_in"], values["check_out"],
                values["break_minutes"], values["total_hours"], values["paid_hours"],
                values["unpaid_hours"], values["overtime_hours"], 0,
                values["unpaid_overtime_amount"], values["unpaid_overtime_paid"],
                values["overtime_payment_date"]
            ))
            message = "Αποθηκεύτηκε"
            self.last_edit_snapshot = None
        else:
            self.cursor.execute("""
                SELECT employee_id, date, check_in, check_out, break_minutes,
                       total_hours, paid_hours, unpaid_hours, overtime_hours,
                       unpaid_overtime_rate, unpaid_overtime_amount, unpaid_overtime_paid,
                       overtime_payment_date
                FROM attendance
                WHERE id = ?
            """, (self.editing_record_id,))

            previous_record = self.cursor.fetchone()

            if previous_record is not None:
                self.last_edit_snapshot = {
                    "id": self.editing_record_id,
                    "values": previous_record
                }

            self.cursor.execute("""
                UPDATE attendance
                SET employee_id = ?, date = ?, check_in = ?, check_out = ?,
                    break_minutes = ?, total_hours = ?, paid_hours = ?,
                    unpaid_hours = ?, overtime_hours = ?,
                    unpaid_overtime_rate = ?, unpaid_overtime_amount = ?, unpaid_overtime_paid = ?,
                    overtime_payment_date = ?
                WHERE id = ?
            """, (
                employee_id, date, values["check_in"], values["check_out"],
                values["break_minutes"], values["total_hours"], values["paid_hours"],
                values["unpaid_hours"], values["overtime_hours"], 0,
                values["unpaid_overtime_amount"], values["unpaid_overtime_paid"],
                values["overtime_payment_date"], self.editing_record_id
            ))
            message = "Ενημερώθηκε"

        self.conn.commit()

        self.filter_employee_select.setCurrentIndex(0)
        monthly_index = self.monthly_employee_select.findData(employee_id)
        if monthly_index >= 0:
            self.monthly_employee_select.blockSignals(True)
            self.monthly_employee_select.setCurrentIndex(monthly_index)
            self.monthly_employee_select.blockSignals(False)

        rate_row = self.cursor.execute(
            "SELECT hourly_rate FROM employees WHERE id=?", (employee_id,)
        ).fetchone()
        hourly_rate = float(rate_row[0] or 0) if rate_row else 0.0
        daily_amount = values["total_hours"] * hourly_rate
        self.result_label.setText(
            f"{message}: {values['total_hours']:.2f} πραγματικές ώρες × "
            f"{hourly_rate:.2f} €/ώρα = {daily_amount:.2f} €"
        )

        self.clear_attendance_form()
        self.load_employees()
        self.load_history()
        self.load_monthly_summary()

    def save_monthly_payroll(self):
        employee_id = self.monthly_employee_select.currentData()
        try:
            year = int(self.monthly_year_input.text().strip())
            month = int(self.monthly_month_select.currentData())
            paid_hours = self.parse_number(self.monthly_paid_hours_input.text())
            rate_text = self.monthly_hourly_rate_input.text().strip()
            hourly_rate = self.parse_number(rate_text) if rate_text else None
            if year < 1900 or year > 9999 or paid_hours < 0 or (hourly_rate is not None and hourly_rate < 0):
                raise ValueError
        except (ValueError, TypeError):
            QMessageBox.warning(
                self, "Σφάλμα",
                "Έλεγξε το έτος και τις ήδη πληρωμένες ώρες. Αν συμπληρώσεις ωρομίσθιο, βάλε μη αρνητικό αριθμό."
            )
            return
        if employee_id is None:
            QMessageBox.warning(self, "Σφάλμα", "Επίλεξε υπάλληλο.")
            return
        if hourly_rate is None:
            hourly_rate = self.hourly_rate_for_date(employee_id, f"{year:04d}-{month:02d}-01")
        now = datetime.now().isoformat(timespec="seconds")
        self.cursor.execute("""
            INSERT INTO monthly_payments(employee_id, year, month, paid_hours, hourly_rate, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(employee_id, year, month) DO UPDATE SET
                paid_hours=excluded.paid_hours,
                hourly_rate=excluded.hourly_rate,
                updated_at=excluded.updated_at
        """, (employee_id, year, month, paid_hours, hourly_rate, now, now))
        self.conn.commit()
        self.load_monthly_summary()
        self.load_history()
        QMessageBox.information(
            self, "Αποθηκεύτηκε",
            f"Καταχωρήθηκαν {paid_hours:.2f} ήδη πληρωμένες ώρες. "
            "Θα καλύψουν πρώτα τις απλήρωτες ώρες των παλαιότερων μηνών. "
            "Τυχόν επιπλέον ώρες μένουν αδιάθετες για μελλοντική εργασία."
        )

    def refresh_monthly_views(self, *_args):
        total_mode = self.payroll_period_mode.currentData() == "total"
        self.monthly_month_select.setVisible(not total_mode)
        self.monthly_year_input.setVisible(not total_mode)
        # The labels share a row with the controls, so switch the matching labels too.
        for column in (2, 3):
            label = self.payroll_selectors.itemAtPosition(0, column).widget()
            label.setVisible(not total_mode)
        self.monthly_adjustments_toggle.setEnabled(
            self.monthly_employee_select.currentData() is not None
        )
        if total_mode:
            self.monthly_adjustments_panel.setVisible(False)
            self.monthly_adjustments_toggle.setChecked(False)
        self.fill_default_monthly_rate()
        self.load_monthly_summary()
        self.load_history()

    def all_payroll_months(self, employee_id):
        months = set()
        for (date_value,) in self.cursor.execute(
            "SELECT date FROM attendance WHERE employee_id=?", (employee_id,)
        ).fetchall():
            try:
                normalized = datetime.strptime(
                    str(date_value), "%Y-%m-%d" if "-" in str(date_value) else "%Y/%m/%d"
                )
                months.add((normalized.year, normalized.month))
            except (TypeError, ValueError):
                pass
        for (start_date, end_date) in self.cursor.execute(
            "SELECT start_date, end_date FROM attendance_periods WHERE employee_id=?", (employee_id,)
        ).fetchall():
            try:
                cursor_date = datetime.strptime(start_date, "%Y-%m-%d").date().replace(day=1)
                last_date = datetime.strptime(end_date, "%Y-%m-%d").date()
                while cursor_date <= last_date:
                    months.add((cursor_date.year, cursor_date.month))
                    cursor_date = (cursor_date.replace(day=28) + timedelta(days=4)).replace(day=1)
            except (TypeError, ValueError):
                pass
        for table in ("monthly_payments", "payroll_payments"):
            for year, month in self.cursor.execute(
                f"SELECT year, month FROM {table} WHERE employee_id=?", (employee_id,)
            ).fetchall():
                months.add((int(year), int(month)))
        today = QDate.currentDate()
        months.add((today.year(), today.month()))
        return sorted(months)

    def calculate_total_payroll(self, employee_id):
        totals = {
            "actual_hours": 0.0, "paid_hours": 0.0, "paid_equivalent_hours": 0.0,
            "cash_payments": 0.0, "total_paid_amount": 0.0, "gross_pay": 0.0,
            "outstanding_hours": 0.0, "extra_paid_hours": 0.0,
            "extra_paid_amount": 0.0, "amount_due": 0.0,
        }
        for year, month in self.all_payroll_months(employee_id):
            result = self.calculate_monthly_payroll(employee_id, year, month)
            for key in totals:
                totals[key] += result[key]
        totals["hourly_rate"] = totals["gross_pay"] / totals["actual_hours"] if totals["actual_hours"] else self.hourly_rate_for_date(
            employee_id, QDate.currentDate().toString("yyyy-MM-dd")
        )
        return totals

    def fill_default_monthly_rate(self):
        if not hasattr(self, "monthly_employee_select") or self.monthly_employee_select.currentData() is None:
            return
        try:
            year = int(self.monthly_year_input.text().strip())
            month = int(self.monthly_month_select.currentData())
        except (ValueError, TypeError):
            return
        employee_id = self.monthly_employee_select.currentData()
        saved = self.cursor.execute(
            "SELECT hourly_rate FROM monthly_payments WHERE employee_id=? AND year=? AND month=?",
            (employee_id, year, month),
        ).fetchone()
        if saved:
            rate = saved[0]
        else:
            rate = self.hourly_rate_for_date(employee_id, f"{year:04d}-{month:02d}-01")
        self.monthly_hourly_rate_input.setText(str(float(rate or 0)))

    def calculate_monthly_payroll(self, employee_id, year, month):
        """Allocate previously paid hours to the oldest unpaid work, without hour overpayments."""
        periods = set(self.all_payroll_months(employee_id))
        periods.add((int(year), int(month)))
        periods = sorted(periods)
        raw = {
            period: self._calculate_monthly_payroll_base(employee_id, *period)
            for period in periods
        }
        prepaid_hours = sum(float(row[0] or 0) for row in self.cursor.execute(
            "SELECT paid_hours FROM monthly_payments WHERE employee_id=?", (employee_id,)
        ).fetchall())

        for period in periods:
            payroll = raw[period]
            rate = payroll["hourly_rate"]
            if rate <= 0 or prepaid_hours <= 0:
                continue
            # Cash payments cover their own month first. Prepaid hours then cover
            # remaining work chronologically, so they cannot create a negative balance.
            cash_due = max(0.0, payroll["gross_pay"] - payroll["cash_payments"])
            unpaid_hours = min(payroll["actual_hours"], cash_due / rate)
            allocated_hours = min(prepaid_hours, unpaid_hours)
            prepaid_hours -= allocated_hours
            payroll["paid_hours"] = allocated_hours
            payroll["total_paid_amount"] += allocated_hours * rate
            payroll["amount_due"] = payroll["gross_pay"] - payroll["total_paid_amount"]
            payroll["extra_paid_amount"] = max(0.0, payroll["total_paid_amount"] - payroll["gross_pay"])
            payroll["paid_equivalent_hours"] = payroll["total_paid_amount"] / rate
            payroll["outstanding_hours"] = max(0.0, payroll["amount_due"] / rate)
            payroll["extra_paid_hours"] = max(0.0, payroll["extra_paid_amount"] / rate)

        return raw[(int(year), int(month))]

    def _calculate_monthly_payroll_base(self, employee_id, year, month):
        start = f"{year:04d}-{month:02d}-01"
        end = f"{year:04d}-{month + 1:02d}-01" if month < 12 else f"{year + 1:04d}-01-01"
        rows = self.cursor.execute(
            "SELECT date, total_hours, check_in, check_out, break_minutes FROM attendance WHERE employee_id=?",
            (employee_id,),
        ).fetchall()
        actual_hours = 0.0
        gross_pay = 0.0
        for date_value, stored_hours, check_in, check_out, break_minutes in rows:
            try:
                normalized_date = datetime.strptime(
                    str(date_value), "%Y-%m-%d" if "-" in str(date_value) else "%Y/%m/%d"
                ).strftime("%Y-%m-%d")
                if not (start <= normalized_date < end):
                    continue
                start_time = datetime.strptime(self.normalize_time(check_in or ""), "%H:%M")
                end_time = datetime.strptime(self.normalize_time(check_out or ""), "%H:%M")
                if end_time <= start_time:
                    from datetime import timedelta
                    end_time += timedelta(days=1)
                hours = max(
                    0.0,
                    (end_time - start_time).total_seconds() / 3600 - int(break_minutes or 0) / 60,
                )
                actual_hours += hours
                gross_pay += hours * self.hourly_rate_for_date(employee_id, normalized_date)
            except (TypeError, ValueError):
                try:
                    normalized_date = datetime.strptime(
                        str(date_value), "%Y-%m-%d" if "-" in str(date_value) else "%Y/%m/%d"
                    ).strftime("%Y-%m-%d")
                    if start <= normalized_date < end:
                        hours = max(0.0, float(stored_hours or 0))
                        actual_hours += hours
                        gross_pay += hours * self.hourly_rate_for_date(employee_id, normalized_date)
                except (TypeError, ValueError):
                    continue

        from datetime import timedelta
        month_end = (datetime.strptime(end, "%Y-%m-%d").date() - timedelta(days=1)).isoformat()
        period_rows = self.cursor.execute("""
            SELECT start_date, end_date, total_work_hours
            FROM attendance_periods
            WHERE employee_id=? AND end_date>=? AND start_date<?
        """, (employee_id, start, end)).fetchall()
        for period_start, period_end, period_hours in period_rows:
            overlap_start = max(period_start, start)
            overlap_end = min(period_end, month_end)
            period_start_date = datetime.strptime(period_start, "%Y-%m-%d").date()
            period_end_date = datetime.strptime(period_end, "%Y-%m-%d").date()
            day_count = (period_end_date - period_start_date).days + 1
            for day_number in range((datetime.strptime(overlap_end, "%Y-%m-%d").date() - datetime.strptime(overlap_start, "%Y-%m-%d").date()).days + 1):
                day = datetime.strptime(overlap_start, "%Y-%m-%d").date() + timedelta(days=day_number)
                hours = float(period_hours or 0) / max(1, day_count)
                actual_hours += hours
                gross_pay += hours * self.hourly_rate_for_date(employee_id, day.isoformat())

        payment = self.cursor.execute(
            "SELECT hourly_rate FROM monthly_payments WHERE employee_id=? AND year=? AND month=?",
            (employee_id, year, month),
        ).fetchone()
        if payment:
            paid_hours, hourly_rate = 0.0, float(payment[0] or 0)
        else:
            paid_hours = 0.0
            rate_row = self.cursor.execute(
                "SELECT hourly_rate FROM employees WHERE id=?", (employee_id,)
            ).fetchone()
            hourly_rate = float(rate_row[0] or 0) if rate_row else 0.0
        cash_payments = float(self.cursor.execute(
            "SELECT COALESCE(SUM(amount), 0) FROM payroll_payments WHERE employee_id=? AND year=? AND month=?",
            (employee_id, year, month),
        ).fetchone()[0] or 0)
        if not payment and actual_hours > 0:
            hourly_rate = gross_pay / actual_hours
        total_paid_amount = paid_hours * hourly_rate + cash_payments
        amount_due = gross_pay - total_paid_amount
        extra_paid_amount = max(0.0, total_paid_amount - gross_pay)
        if hourly_rate > 0:
            paid_equivalent_hours = total_paid_amount / hourly_rate
            outstanding_hours = amount_due / hourly_rate
            extra_paid_hours = extra_paid_amount / hourly_rate
        else:
            paid_equivalent_hours = paid_hours
            outstanding_hours = max(0.0, actual_hours - paid_hours)
            extra_paid_hours = max(0.0, paid_hours - actual_hours)
        return {
            "actual_hours": actual_hours,
            "paid_hours": paid_hours,
            "paid_equivalent_hours": paid_equivalent_hours,
            "cash_payments": cash_payments,
            "total_paid_amount": total_paid_amount,
            "gross_pay": gross_pay,
            "outstanding_hours": outstanding_hours,
            "extra_paid_hours": extra_paid_hours,
            "extra_paid_amount": extra_paid_amount,
            "hourly_rate": hourly_rate,
            "amount_due": amount_due,
        }

    def load_monthly_summary(self):
        if not hasattr(self, "monthly_summary_label"):
            return
        employee_id = self.monthly_employee_select.currentData()
        total_mode = self.payroll_period_mode.currentData() == "total"
        self.add_payment_btn.setEnabled(employee_id is not None)
        self.monthly_adjustments_toggle.setEnabled(employee_id is not None)
        self.load_payment_history()
        try:
            year = int(self.monthly_year_input.text().strip())
            month = int(self.monthly_month_select.currentData())
        except (ValueError, TypeError):
            self.monthly_summary_label.setText("Έλεγξε μήνα και έτος.")
            return
        if employee_id is None:
            self.monthly_total_value.setText("—")
            self.monthly_paid_value.setText("—")
            self.monthly_due_value.setText("—")
            self.monthly_summary_label.setText(
                "Ο πίνακας δείχνει όλους τους υπαλλήλους. Επίλεξε έναν για ατομική σύνοψη και πληρωμή."
            )
            return
        payroll = self.calculate_total_payroll(employee_id) if total_mode else self.calculate_monthly_payroll(employee_id, year, month)
        saved_hours = self.cursor.execute(
            "SELECT paid_hours, hourly_rate FROM monthly_payments WHERE employee_id=? AND year=? AND month=?",
            (employee_id, year, month),
        ).fetchone()
        self.monthly_paid_hours_input.setText(str(float(saved_hours[0] or 0)) if saved_hours else "0")
        self.monthly_hourly_rate_input.setText(
            str(float(saved_hours[1] or 0)) if saved_hours
            else str(self.hourly_rate_for_date(employee_id, f"{year:04d}-{month:02d}-01"))
        )
        self.monthly_total_value.setText(f"{payroll['gross_pay']:.2f} €")
        self.monthly_paid_value.setText(f"{payroll['total_paid_amount']:.2f} €")
        self.monthly_due_value.setText(f"{payroll['amount_due']:.2f} €")
        self.monthly_summary_label.setText(
            f"{payroll['actual_hours']:.2f} πραγματικές ώρες · "
            f"{payroll['paid_equivalent_hours']:.2f} πληρωμένες ώρες · "
            f"{payroll['outstanding_hours']:.2f} ώρες υπόλοιπο · "
            f"ωρομίσθιο {payroll['hourly_rate']:.2f} € · "
            f"Υπόλοιπο: {payroll['amount_due']:.2f} €" +
            (" (υπερπληρωμή — οφείλει στην επιχείρηση)" if payroll['amount_due'] < 0 else "")
        )

    def add_monthly_payment(self):
        employee_id = self.monthly_employee_select.currentData()
        if employee_id is None:
            QMessageBox.warning(self, "Επίλεξε υπάλληλο", "Επίλεξε έναν συγκεκριμένο υπάλληλο πριν από την πληρωμή.")
            return

        try:
            amount = self.parse_number(self.new_payment_amount_input.text())
            if amount <= 0:
                raise ValueError
        except (ValueError, TypeError):
            QMessageBox.warning(self, "Σφάλμα", "Βάλε ποσό πληρωμής μεγαλύτερο από 0.")
            return

        periods = [
            (year, month, self.calculate_monthly_payroll(employee_id, year, month))
            for year, month in self.all_payroll_months(employee_id)
        ]
        if not periods:
            today = QDate.currentDate()
            periods = [(today.year(), today.month(),
                        self.calculate_monthly_payroll(employee_id, today.year(), today.month()))]
        amount_cents = int(round(amount * 100))
        if amount_cents <= 0:
            QMessageBox.warning(self, "Σφάλμα", "Το ποσό πρέπει να είναι τουλάχιστον 0,01 €.")
            return
        allocations = [0] * len(periods)
        due_cents = [max(0, int(round(row[2]["amount_due"] * 100))) for row in periods]
        remaining_cents = amount_cents
        for index, due in enumerate(due_cents):
            if remaining_cents <= 0:
                break
            allocations[index] = min(remaining_cents, due)
            remaining_cents -= allocations[index]
        if remaining_cents > 0:
            allocations[-1] += remaining_cents

        month_names = ("Ιανουάριος", "Φεβρουάριος", "Μάρτιος", "Απρίλιος", "Μάιος", "Ιούνιος",
                       "Ιούλιος", "Αύγουστος", "Σεπτέμβριος", "Οκτώβριος", "Νοέμβριος", "Δεκέμβριος")
        distribution = [
            f"{month_names[month - 1]} {year}: {allocated / 100:.2f} €"
            for (year, month, _payroll), allocated in zip(periods, allocations) if allocated > 0
        ]
        preview = "\n".join(distribution) or ""
        if not preview:
            preview = f"{periods[-1][1]:02d}/{periods[-1][0]}: {amount_cents / 100:.2f} € (προκαταβολή)"
        confirmation = QMessageBox.question(
            self,
            "Επιβεβαίωση πληρωμής",
            f"Καταχώριση πληρωμής {amount_cents / 100:.2f} €;\n\n"
            f"{preview}\n\n"
            "Η πληρωμή καλύπτει πρώτα τους παλαιότερους οφειλόμενους μήνες. "
            "Τυχόν επιπλέον ποσό προστίθεται στον πιο πρόσφατο μήνα και εμφανίζεται ως αρνητικό υπόλοιπο.",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No,
        )
        if confirmation != QMessageBox.Yes:
            return

        created_at = datetime.now().isoformat(timespec="seconds")
        try:
            self.cursor.execute("BEGIN IMMEDIATE")
            for (year, month, _payroll), allocated in zip(periods, allocations):
                if allocated <= 0:
                    continue
                self.cursor.execute(
                    "INSERT INTO payroll_payments(employee_id, year, month, amount, created_at) VALUES (?, ?, ?, ?, ?)",
                    (employee_id, year, month, allocated / 100, created_at),
                )
            self.conn.commit()
        except sqlite3.Error as error:
            self.conn.rollback()
            QMessageBox.critical(self, "Σφάλμα", f"Η πληρωμή δεν αποθηκεύτηκε: {error}")
            return

        self.new_payment_amount_input.clear()
        last_paid_index = max(
            (index for index, allocated in enumerate(allocations) if allocated > 0),
            default=len(periods) - 1,
        )
        view_year, view_month, _ = periods[last_paid_index]
        self.monthly_month_select.setCurrentIndex(view_month - 1)
        self.monthly_year_input.setText(str(view_year))
        self.load_monthly_summary()
        self.load_history()

    def clear_attendance_form(self):
        self.editing_record_id = None
        self.editing_period_id = None
        self.save_btn.setText("Αποθήκευση Ωρών")
        self.check_in_input.clear()
        self.check_out_input.clear()
        self.break_input.clear()
        if hasattr(self, "aggregate_work_hours_input"):
            self.aggregate_work_hours_input.clear()
            self.aggregate_break_hours_input.clear()
            self.aggregate_attendance_save_btn.setText("Αποθήκευση Συγκεντρωτικών Ωρών")

    def toggle_aggregate_attendance(self, expanded):
        self.aggregate_attendance_panel.setVisible(expanded)
        self.aggregate_attendance_toggle.setArrowType(Qt.DownArrow if expanded else Qt.RightArrow)

    def save_aggregate_attendance(self):
        employee_id = self.employee_select.currentData()
        if employee_id is None:
            QMessageBox.warning(self, "Επίλεξε υπάλληλο", "Επίλεξε υπάλληλο πριν αποθηκεύσεις τις ώρες.")
            return
        start_date = self.aggregate_start_date_input.date().toString("yyyy-MM-dd")
        end_date = self.aggregate_end_date_input.date().toString("yyyy-MM-dd")
        if end_date < start_date:
            QMessageBox.warning(self, "Μη έγκυρο διάστημα", "Η ημερομηνία λήξης πρέπει να είναι μετά την ημερομηνία έναρξης.")
            return
        try:
            work_hours = self.parse_number(self.aggregate_work_hours_input.text())
            break_hours = self.parse_number(self.aggregate_break_hours_input.text())
            if work_hours < 0 or break_hours < 0 or (work_hours == 0 and break_hours == 0):
                raise ValueError
        except ValueError:
            QMessageBox.warning(
                self,
                "Μη έγκυρες ώρες",
                "Βάλε μη αρνητικές καθαρές ώρες και ώρες διαλείμματος, με τουλάχιστον μία τιμή μεγαλύτερη από 0.",
            )
            return

        if self.editing_period_id is None:
            self.cursor.execute("""
                INSERT INTO attendance_periods(
                    employee_id, start_date, end_date, total_work_hours, break_hours, created_at
                ) VALUES (?, ?, ?, ?, ?, ?)
            """, (employee_id, start_date, end_date, work_hours, break_hours,
                  datetime.now().isoformat(timespec="seconds")))
            action = "Καταχωρήθηκαν"
        else:
            self.cursor.execute("""
                UPDATE attendance_periods
                SET employee_id=?, start_date=?, end_date=?, total_work_hours=?, break_hours=?
                WHERE id=?
            """, (employee_id, start_date, end_date, work_hours, break_hours, self.editing_period_id))
            action = "Ενημερώθηκαν"
        self.conn.commit()
        self.editing_period_id = None
        self.aggregate_attendance_save_btn.setText("Αποθήκευση Συγκεντρωτικών Ωρών")
        self.aggregate_work_hours_input.clear()
        self.aggregate_break_hours_input.clear()
        self.result_label.setText(
            f"{action} {work_hours:.2f} καθαρές ώρες και {break_hours:.2f} ώρες διαλείμματος "
            f"για {start_date} έως {end_date}."
        )
        self.load_history()
        self.load_monthly_summary()

    @staticmethod
    def period_hours_in_range(period_start, period_end, total_hours, range_start, range_end):
        try:
            period_start = datetime.strptime(period_start, "%Y-%m-%d").date()
            period_end = datetime.strptime(period_end, "%Y-%m-%d").date()
            range_start = datetime.strptime(range_start, "%Y-%m-%d").date()
            range_end = datetime.strptime(range_end, "%Y-%m-%d").date()
        except (TypeError, ValueError):
            return 0.0
        overlap_start = max(period_start, range_start)
        overlap_end = min(period_end, range_end)
        if overlap_end < overlap_start:
            return 0.0
        total_days = (period_end - period_start).days + 1
        overlap_days = (overlap_end - overlap_start).days + 1
        if total_days <= 0:
            return 0.0
        return float(total_hours or 0) * overlap_days / total_days

    def cancel_edit(self):
        self.clear_attendance_form()
        self.result_label.setText("Η επεξεργασία ακυρώθηκε")

    def clear_filters(self):
        self.filter_employee_select.setCurrentIndex(0)
        self.from_date_input.setDate(self.get_data_start_date())
        self.to_date_input.setDate(QDate.currentDate())
        self.load_history()

    def get_selected_record_id(self):
        selected_row = self.history_table.currentRow()

        if selected_row < 0:
            QMessageBox.warning(self, "Σφάλμα", "Επίλεξε μία καταχώρηση")
            return None

        item = self.history_table.item(selected_row, 0)

        if item is None:
            return None

        record_id = item.data(Qt.UserRole)
        return int(record_id if record_id is not None else item.text())

    def load_selected_record_for_edit(self):
        record_id = self.get_selected_record_id()
        if record_id is None:
            return

        if record_id < 0:
            self.cursor.execute("""
                SELECT employee_id, start_date, end_date, total_work_hours, break_hours
                FROM attendance_periods WHERE id=?
            """, (-record_id,))
            period = self.cursor.fetchone()
            if period is None:
                QMessageBox.warning(self, "Σφάλμα", "Δεν βρέθηκε η συγκεντρωτική καταχώριση.")
                return
            employee_id, start_date, end_date, work_hours, break_hours = period
            employee_index = self.employee_select.findData(employee_id)
            if employee_index >= 0:
                self.employee_select.setCurrentIndex(employee_index)
            start_value = QDate.fromString(start_date, "yyyy-MM-dd")
            end_value = QDate.fromString(end_date, "yyyy-MM-dd")
            if start_value.isValid():
                self.aggregate_start_date_input.setDate(start_value)
            if end_value.isValid():
                self.aggregate_end_date_input.setDate(end_value)
            self.aggregate_work_hours_input.setText(str(work_hours))
            self.aggregate_break_hours_input.setText(str(break_hours))
            self.editing_period_id = -record_id
            self.aggregate_attendance_toggle.setChecked(True)
            self.aggregate_attendance_save_btn.setText("Αποθήκευση Αλλαγών Συγκεντρωτικών Ωρών")
            self.result_label.setText("Επεξεργασία συγκεντρωτικής καταχώρισης")
            if hasattr(self, "tabs"):
                self.tabs.setCurrentIndex(1)
            return

        self.cursor.execute("""
            SELECT id, employee_id, date, check_in, check_out, break_minutes
            FROM attendance WHERE id = ?
        """, (record_id,))

        record = self.cursor.fetchone()

        if record is None:
            QMessageBox.warning(self, "Σφάλμα", "Δεν βρέθηκε η καταχώρηση")
            return

        (
            attendance_id, employee_id, date, check_in, check_out, break_minutes
        ) = record

        self.editing_record_id = attendance_id

        index = self.employee_select.findData(employee_id)
        if index >= 0:
            self.employee_select.setCurrentIndex(index)

        parsed_date = QDate.fromString(date, "yyyy-MM-dd")
        if parsed_date.isValid():
            self.date_input.setDate(parsed_date)

        self.check_in_input.setText(check_in)
        self.check_out_input.setText(check_out)
        self.break_input.setText(str(break_minutes or 0))

        self.save_btn.setText("Ενημέρωση Καταχώρησης")
        self.result_label.setText(f"Επεξεργασία ID {attendance_id}")

        if hasattr(self, "tabs"):
            self.tabs.setCurrentIndex(1)

    def undo_last_edit(self):
        if not self.last_edit_snapshot:
            QMessageBox.information(self, "Αναίρεση", "Δεν υπάρχει πρόσφατη αλλαγή μισθοδοσίας για αναίρεση.")
            return

        record_id = self.last_edit_snapshot["id"]
        previous_values = self.last_edit_snapshot["values"]

        self.cursor.execute("""
            UPDATE attendance
            SET employee_id = ?,
                date = ?,
                check_in = ?,
                check_out = ?,
                break_minutes = ?,
                total_hours = ?,
                paid_hours = ?,
                unpaid_hours = ?,
                overtime_hours = ?,
                unpaid_overtime_rate = ?,
                unpaid_overtime_amount = ?,
                unpaid_overtime_paid = ?,
                overtime_payment_date = ?
            WHERE id = ?
        """, (*previous_values, record_id))

        self.conn.commit()
        self.last_edit_snapshot = None
        self.clear_attendance_form()
        self.load_employees()
        self.load_history()

        QMessageBox.information(self, "Αναίρεση", f"Η αλλαγή της καταχώρησης ID {record_id} αναιρέθηκε.")

    def delete_selected_record(self):
        record_id = self.get_selected_record_id()
        if record_id is None:
            return

        record_label = "συγκεντρωτική καταχώριση" if record_id < 0 else f"καταχώριση ID {record_id}"
        confirm = QMessageBox.question(self, "Επιβεβαίωση", f"Να διαγραφεί η {record_label};")
        if confirm != QMessageBox.Yes:
            return

        if record_id < 0:
            self.cursor.execute("DELETE FROM attendance_periods WHERE id = ?", (-record_id,))
        else:
            self.cursor.execute("DELETE FROM attendance WHERE id = ?", (record_id,))
        self.conn.commit()
        self.load_history()

    def get_filtered_records(self, employee_id=None, from_date=None, to_date=None):
        report_filters = from_date is not None or to_date is not None
        search = ""
        if not report_filters:
            employee_id = self.filter_employee_select.currentData()
            if employee_id is None and self.filter_employee_select.currentIndex() != 0:
                search = self.filter_employee_select.currentText().strip()
            from_date = self.from_date_input.date().toString("yyyy-MM-dd")
            to_date = self.to_date_input.date().toString("yyyy-MM-dd")
        else:
            from_date = from_date or self.report_from_date.date().toString("yyyy-MM-dd")
            to_date = to_date or self.report_to_date.date().toString("yyyy-MM-dd")

        attendance_query = """
            SELECT attendance.id, employees.name, attendance.date,
                   attendance.check_in, attendance.check_out,
                   attendance.break_minutes, attendance.total_hours
            FROM attendance
            JOIN employees ON attendance.employee_id = employees.id
            WHERE attendance.date >= ?
            AND attendance.date <= ?
        """

        attendance_params = [from_date, to_date]
        period_params = [to_date, from_date]

        if employee_id is not None:
            attendance_query += " AND attendance.employee_id = ?"
            attendance_params.append(employee_id)
            period_filter = " AND periods.employee_id = ?"
            period_params.append(employee_id)
        elif search and search != "Όλοι οι υπάλληλοι":
            attendance_query += " AND employees.name LIKE ?"
            attendance_params.append(f"%{search}%")
            period_filter = " AND employees.name LIKE ?"
            period_params.append(f"%{search}%")
        else:
            period_filter = ""

        self.cursor.execute(attendance_query, attendance_params)
        rows = [(row[2], row) for row in self.cursor.fetchall()]
        period_query = f"""
            SELECT -periods.id, employees.name,
                   periods.start_date || ' – ' || periods.end_date,
                   'Συγκεντρωτικά', '', periods.break_hours,
                   periods.total_work_hours, periods.start_date, periods.end_date
            FROM attendance_periods AS periods
            JOIN employees ON employees.id = periods.employee_id
            WHERE periods.start_date <= ? AND periods.end_date >= ? {period_filter}
        """
        self.cursor.execute(period_query, period_params)
        for row in self.cursor.fetchall():
            period_id, name, label, check_in, check_out, break_hours, work_hours, start, end = row
            clipped_work_hours = self.period_hours_in_range(
                start, end, work_hours, from_date, to_date
            )
            clipped_break_hours = self.period_hours_in_range(
                start, end, break_hours, from_date, to_date
            )
            rows.append((end, (
                period_id, name, label, check_in, check_out,
                int(round(clipped_break_hours * 60)), clipped_work_hours,
            )))

        def sort_date(value):
            try:
                return datetime.strptime(value, "%Y-%m-%d")
            except (TypeError, ValueError):
                try:
                    return datetime.strptime(value, "%Y/%m/%d")
                except (TypeError, ValueError):
                    return datetime.min

        rows.sort(key=lambda item: (sort_date(item[0]), item[1][0]), reverse=True)
        return [row for _, row in rows]

    def load_history(self):
        records = self.get_filtered_records()

        self.history_table.clearContents()
        self.history_table.setRowCount(len(records))
        total_hours_sum = 0.0

        for row, record in enumerate(records):
            total_hours_sum += float(record[6] or 0)
            for col, value in enumerate(record):
                display_value = "Περίοδος" if col == 0 and int(record[0]) < 0 else str(value)
                item = QTableWidgetItem(display_value)
                if col == 0:
                    item.setData(Qt.UserRole, int(record[0]))
                self.history_table.setItem(row, col, item)
        self.history_table.resizeColumnsToContents()
        self.summary_label.setText(f"Ώρες παρουσίας στο επιλεγμένο φίλτρο: {total_hours_sum:.2f}")
        if hasattr(self, "report_table"):
            self.refresh_report()

        total_mode = self.payroll_period_mode.currentData() == "total"
        if not total_mode:
            try:
                month = int(self.monthly_month_select.currentData())
                year = int(self.monthly_year_input.text().strip())
            except (ValueError, TypeError):
                self.employee_summary_table.clearContents()
                self.employee_summary_table.setRowCount(0)
                return
        self.employee_summary_group.setTitle(
            "Συνολική μισθοδοσία" if total_mode else "Μισθοδοσία ανά μήνα"
        )
        if not total_mode:
            self.load_monthly_breakdown(year)
            return
        self.employee_summary_table.setHorizontalHeaderLabels([
            "Υπάλληλος", "Πραγματικές Ώρες", "Ώρες ήδη πληρωμένες",
            "Υπόλοιπο Ωρών", "Σύνολο Μισθοδοσίας", "Πληρωμένο Ποσό",
            "Υπόλοιπο (€) · αρνητικό = υπερπληρωμή"
        ])
        self.employee_summary_table.setColumnCount(7)

        self.employee_summary_table.clearContents()
        selected_employee_id = self.monthly_employee_select.currentData()
        if selected_employee_id is None:
            employees = self.cursor.execute("SELECT id, name FROM employees ORDER BY name").fetchall()
        else:
            employee = self.cursor.execute(
                "SELECT id, name FROM employees WHERE id=?", (selected_employee_id,)
            ).fetchone()
            employees = [employee] if employee else []

        summary_rows = []
        for employee_id, employee in employees:
            payroll = self.calculate_total_payroll(employee_id) if total_mode else self.calculate_monthly_payroll(employee_id, year, month)
            if selected_employee_id is None and not any((
                payroll["actual_hours"], payroll["paid_hours"], payroll["total_paid_amount"]
            )):
                continue
            summary_rows.append((employee_id, employee, payroll))

        self.employee_summary_table.setRowCount(len(summary_rows))
        for row, (employee_id, employee, payroll) in enumerate(summary_rows):
            month_hours = payroll["actual_hours"]
            row_values = [employee, f"{month_hours:.2f}", f"{payroll['paid_hours']:.2f}",
                          f"{payroll['outstanding_hours']:.2f}" + (f" (επιπλέον {payroll['extra_paid_hours']:.2f})" if payroll['extra_paid_hours'] else ""),
                          f"{payroll['gross_pay']:.2f} €", f"{payroll['total_paid_amount']:.2f} €",
                          f"{payroll['amount_due']:.2f} €"]

            for col, value in enumerate(row_values):
                item = QTableWidgetItem(value)
                if col == 0:
                    item.setData(Qt.UserRole, employee_id)
                self.employee_summary_table.setItem(row, col, item)

        self.employee_summary_table.resizeColumnsToContents()

    def load_monthly_breakdown(self, year):
        table = self.employee_summary_table
        table.setColumnCount(9)
        table.setHorizontalHeaderLabels([
            "Υπάλληλος", "Μήνας", "Πραγματικές ώρες", "Ώρες ήδη πληρωμένες", "Σύνολο μισθοδοσίας",
            "Πληρωμές που κάλυψαν τον μήνα", "Σύνολο πληρωμών", "Κατάσταση", "Υπόλοιπο"
        ])
        table.clearContents()
        selected_employee_id = self.monthly_employee_select.currentData()
        if selected_employee_id is None:
            employees = self.cursor.execute("SELECT id, name FROM employees ORDER BY name").fetchall()
        else:
            row = self.cursor.execute(
                "SELECT id, name FROM employees WHERE id=?", (selected_employee_id,)
            ).fetchone()
            employees = [row] if row else []

        names = ("Ιανουάριος", "Φεβρουάριος", "Μάρτιος", "Απρίλιος", "Μάιος", "Ιούνιος",
                 "Ιούλιος", "Αύγουστος", "Σεπτέμβριος", "Οκτώβριος", "Νοέμβριος", "Δεκέμβριος")
        rows = []
        for employee_id, employee_name in employees:
            for month in range(1, 13):
                payroll = self.calculate_monthly_payroll(employee_id, year, month)
                if not any((payroll["actual_hours"], payroll["paid_hours"],
                            payroll["total_paid_amount"], payroll["amount_due"])):
                    continue
                payments = self.cursor.execute("""
                    SELECT created_at, amount FROM payroll_payments
                    WHERE employee_id=? AND year=? AND month=?
                    ORDER BY created_at, id
                """, (employee_id, year, month)).fetchall()
                payment_details = []
                for created_at, amount in payments:
                    try:
                        date_label = datetime.fromisoformat(created_at).strftime("%d/%m/%Y")
                    except (TypeError, ValueError):
                        date_label = str(created_at or "")
                    payment_details.append(f"{date_label}: {float(amount or 0):.2f} €")
                hour_pay_amount = float(payroll["paid_hours"] or 0) * float(payroll["hourly_rate"] or 0)
                if hour_pay_amount:
                    payment_details.insert(0, f"Καταχώριση {payroll['paid_hours']:.2f} ωρών ήδη πληρωμένων · αξία {hour_pay_amount:.2f} €")
                if not payment_details:
                    payment_details = ["Δεν έχει καλυφθεί"]
                if payroll["amount_due"] < -0.005:
                    status = "Υπερπληρωμένος"
                elif payroll["amount_due"] <= 0.005:
                    status = "Καλύφθηκε"
                elif payroll["total_paid_amount"] > 0:
                    status = "Μερική κάλυψη"
                else:
                    status = "Ακάλυπτος"
                rows.append((employee_id, employee_name, month, payroll, "\n".join(payment_details), status))

        table.setRowCount(len(rows))
        for row_index, (employee_id, employee_name, month, payroll, details, status) in enumerate(rows):
            values = [
                employee_name, names[month - 1], f"{payroll['actual_hours']:.2f}", f"{payroll['paid_hours']:.2f}",
                f"{payroll['gross_pay']:.2f} €", details,
                f"{payroll['total_paid_amount']:.2f} €", status, f"{payroll['amount_due']:.2f} €",
            ]
            for column, value in enumerate(values):
                item = QTableWidgetItem(value)
                if column == 0:
                    item.setData(Qt.UserRole, employee_id)
                elif column == 1:
                    item.setData(Qt.UserRole, month)
                table.setItem(row_index, column, item)
            table.setRowHeight(row_index, max(36, 22 * len(details.splitlines())))
        table.resizeColumnsToContents()
        table.horizontalHeader().setStretchLastSection(True)

    def export_to_excel(self):
        headers = getattr(self, "current_report_headers", [])
        rows = getattr(self, "current_report_rows", [])
        if not rows:
            QMessageBox.information(self, "Κενή αναφορά", "Δεν υπάρχουν δεδομένα για εξαγωγή με τα τρέχοντα φίλτρα.")
            return

        report_name = self.report_type_select.currentText()
        safe_name = "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in report_name).strip()
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Εξαγωγή αναφοράς σε Excel", f"{safe_name}.xlsx", "Excel Files (*.xlsx)"
        )

        if not file_path:
            return

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = report_name[:31]
        sheet.append(["Αναφορά", report_name])
        sheet.append(["Περίοδος", self.report_from_date.date().toString("dd/MM/yyyy"),
                      self.report_to_date.date().toString("dd/MM/yyyy")])
        sheet.append([])
        sheet.append(headers)
        for row in rows:
            sheet.append(list(row))
        header_row = 4
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
        for column_cells in sheet.columns:
            width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 48)
            sheet.column_dimensions[column_cells[0].column_letter].width = width

        workbook.save(file_path)
        QMessageBox.information(self, "Ολοκληρώθηκε", "Η αναφορά αποθηκεύτηκε σε Excel.")

    def export_employee_payroll_to_excel(self):
        employee_id = self.report_employee_select.currentData()
        if employee_id is None:
            QMessageBox.warning(self, "Επίλεξε υπάλληλο", "Στα φίλτρα της αναφοράς επίλεξε έναν συγκεκριμένο υπάλληλο.")
            return

        employee = self.cursor.execute(
            "SELECT name, hourly_rate FROM employees WHERE id=?", (employee_id,)
        ).fetchone()
        if employee is None:
            QMessageBox.warning(self, "Σφάλμα", "Δεν βρέθηκε ο επιλεγμένος υπάλληλος.")
            return
        employee_name, base_rate = employee
        from_date = self.report_from_date.date().toString("yyyy-MM-dd")
        to_date = self.report_to_date.date().toString("yyyy-MM-dd")
        start_year, start_month = map(int, from_date[:7].split("-"))
        end_year, end_month = map(int, to_date[:7].split("-"))
        start_period = start_year * 100 + start_month
        end_period = end_year * 100 + end_month

        attendance = self.cursor.execute("""
            SELECT date, check_in, check_out, break_minutes, total_hours
            FROM attendance
            WHERE employee_id=? AND date>=? AND date<=?
            ORDER BY date, id
        """, (employee_id, from_date, to_date)).fetchall()
        period_attendance = self.cursor.execute("""
            SELECT start_date, end_date, total_work_hours, break_hours
            FROM attendance_periods
            WHERE employee_id=? AND start_date<=? AND end_date>=?
            ORDER BY start_date, id
        """, (employee_id, to_date, from_date)).fetchall()
        month_payrolls = self.cursor.execute("""
            SELECT year, month, paid_hours, hourly_rate, updated_at
            FROM monthly_payments
            WHERE employee_id=? AND year*100+month BETWEEN ? AND ?
            ORDER BY year, month
        """, (employee_id, start_period, end_period)).fetchall()
        cash_payments = self.cursor.execute("""
            SELECT year, month, amount, created_at
            FROM payroll_payments
            WHERE employee_id=? AND year*100+month BETWEEN ? AND ?
            ORDER BY year, month, created_at, id
        """, (employee_id, start_period, end_period)).fetchall()

        safe_name = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in employee_name).strip("_")
        suggested = f"{safe_name}_payroll_{from_date}_to_{to_date}.xlsx"
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Εξαγωγή Παρουσιών & Πληρωμών", suggested, "Excel Files (*.xlsx)"
        )
        if not file_path:
            return

        workbook = Workbook()
        summary = workbook.active
        summary.title = "Σύνοψη"
        summary.append(["Υπάλληλος", employee_name])
        summary.append(["Περίοδος παρουσιών", from_date, to_date])
        period_hours_in_export = sum(
            self.period_hours_in_range(start, end, hours, from_date, to_date)
            for start, end, hours, _break_hours in period_attendance
        )
        summary.append(["Πραγματικές ώρες", sum(float(row[4] or 0) for row in attendance) + period_hours_in_export])
        summary.append(["Πληρωμένες ώρες payroll", sum(float(row[2] or 0) for row in month_payrolls)])
        monthly_paid_amount = sum(float(row[2] or 0) * float(row[3] or 0) for row in month_payrolls)
        cash_paid_amount = sum(float(row[2] or 0) for row in cash_payments)
        summary.append(["Ποσό πληρωμένων ωρών", monthly_paid_amount])
        summary.append(["Καταχωρημένες πληρωμές ποσού", cash_paid_amount])
        summary.append(["Συνολικά καταχωρημένο πληρωμένο ποσό", monthly_paid_amount + cash_paid_amount])
        summary.append(["Βασικό ωρομίσθιο", float(base_rate or 0)])

        attendance_sheet = workbook.create_sheet("Παρουσίες")
        attendance_sheet.append(["Ημερομηνία", "Προσέλευση", "Αποχώρηση", "Διάλειμμα (λεπτά)", "Πραγματικές ώρες", "Ωρομίσθιο (€)", "Αξία ωρών (€)"])
        gross_amount = 0.0
        for date_value, check_in, check_out, break_minutes, hours in attendance:
            year, month = map(int, date_value[:7].split("-"))
            rate_row = self.cursor.execute(
                "SELECT hourly_rate FROM monthly_payments WHERE employee_id=? AND year=? AND month=?",
                (employee_id, year, month),
            ).fetchone()
            rate = float(rate_row[0] or 0) if rate_row else float(base_rate or 0)
            hours = float(hours or 0)
            gross_amount += hours * rate
            attendance_sheet.append([date_value, check_in, check_out, break_minutes, hours, rate, hours * rate])
        from datetime import timedelta
        for period_start, period_end, total_work_hours, break_hours in period_attendance:
            visible_start = max(period_start, from_date)
            visible_end = min(period_end, to_date)
            current_month = datetime.strptime(visible_start, "%Y-%m-%d").date().replace(day=1)
            last_date = datetime.strptime(visible_end, "%Y-%m-%d").date()
            while current_month <= last_date:
                next_month = (current_month.replace(day=28) + timedelta(days=4)).replace(day=1)
                slice_start = max(datetime.strptime(period_start, "%Y-%m-%d").date(),
                                  datetime.strptime(visible_start, "%Y-%m-%d").date(), current_month)
                slice_end = min(datetime.strptime(period_end, "%Y-%m-%d").date(),
                                last_date, next_month - timedelta(days=1))
                slice_start_text = slice_start.isoformat()
                slice_end_text = slice_end.isoformat()
                hours = self.period_hours_in_range(
                    period_start, period_end, total_work_hours, slice_start_text, slice_end_text
                )
                rest_hours = self.period_hours_in_range(
                    period_start, period_end, break_hours, slice_start_text, slice_end_text
                )
                rate_row = self.cursor.execute(
                    "SELECT hourly_rate FROM monthly_payments WHERE employee_id=? AND year=? AND month=?",
                    (employee_id, current_month.year, current_month.month),
                ).fetchone()
                rate = float(rate_row[0] or 0) if rate_row else float(base_rate or 0)
                gross_amount += hours * rate
                attendance_sheet.append([
                    f"{slice_start_text} έως {slice_end_text} (συγκεντρωτικά)",
                    "Συγκεντρωτικά", "", rest_hours * 60, hours, rate, hours * rate,
                ])
                current_month = next_month
        summary.append(["Αξία παρουσιών στην περίοδο", gross_amount])

        payments_sheet = workbook.create_sheet("Πληρωμές")
        payments_sheet.append(["Έτος", "Μήνας", "Τύπος", "Ποσό (€)", "Ώρες", "Ωρομίσθιο (€)", "Ημερομηνία καταχώρισης"])
        for year, month, paid_hours, rate, updated_at in month_payrolls:
            payments_sheet.append([year, month, "Ώρες ήδη πληρωμένες", float(paid_hours or 0) * float(rate or 0), float(paid_hours or 0), float(rate or 0), updated_at])
        for year, month, amount, created_at in cash_payments:
            payments_sheet.append([year, month, "Μερική πληρωμή", float(amount), "", "", created_at])

        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions if sheet.max_row > 1 else None
            for column_cells in sheet.columns:
                width = min(max(max(len(str(cell.value or "")) for cell in column_cells) + 2, 12), 42)
                sheet.column_dimensions[column_cells[0].column_letter].width = width

        workbook.save(file_path)
        QMessageBox.information(self, "Ολοκληρώθηκε", f"Η εξαγωγή για τον/την {employee_name} αποθηκεύτηκε.")

    def export_to_pdf(self):
        headers = getattr(self, "current_report_headers", [])
        rows = getattr(self, "current_report_rows", [])
        if not rows:
            QMessageBox.information(self, "Κενή αναφορά", "Δεν υπάρχουν δεδομένα για εξαγωγή με τα τρέχοντα φίλτρα.")
            return

        report_name = self.report_type_select.currentText()
        safe_name = "".join(ch if ch.isalnum() or ch in "-_ " else "_" for ch in report_name).strip()
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Εξαγωγή αναφοράς σε PDF", f"{safe_name}.pdf", "PDF Files (*.pdf)"
        )

        if not file_path:
            return

        try:
            font_path = mpl_font_manager.findfont("DejaVu Sans")
            if "PayrollReportUnicode" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("PayrollReportUnicode", font_path))
            regular_font = "PayrollReportUnicode"
        except Exception:
            regular_font = "Helvetica"
        small_style = ParagraphStyle(
            "ReportCell", fontName=regular_font, fontSize=7.2, leading=9, wordWrap="CJK"
        )
        title_style = ParagraphStyle(
            "ReportTitle", fontName=regular_font, fontSize=14, leading=18, spaceAfter=5
        )
        meta_style = ParagraphStyle("ReportMeta", fontName=regular_font, fontSize=8, leading=11)
        document = SimpleDocTemplate(
            file_path, pagesize=landscape(A4), leftMargin=24, rightMargin=24,
            topMargin=26, bottomMargin=24,
        )
        story = [
            Paragraph(escape(report_name), title_style),
            Paragraph(
                f"Περίοδος: {self.report_from_date.date().toString('dd/MM/yyyy')} – "
                f"{self.report_to_date.date().toString('dd/MM/yyyy')}", meta_style
            ),
            Paragraph(escape(self.report_summary_label.text()), meta_style),
            Spacer(1, 10),
        ]
        def cell(value):
            return Paragraph(escape(str(value)).replace("\n", "<br/>"), small_style)
        data = [[cell(value) for value in headers]]
        data.extend([[cell(value) for value in row] for row in rows])
        available_width = landscape(A4)[0] - 48
        column_width = available_width / max(1, len(headers))
        report_table = Table(data, colWidths=[column_width] * len(headers), repeatRows=1)
        report_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3978d4")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#aab0b8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#edf1f5")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(report_table)
        document.build(story)
        QMessageBox.information(self, "Ολοκληρώθηκε", "Η αναφορά αποθηκεύτηκε σε PDF.")

    def backup_database(self):
        if not os.path.exists(DB_NAME):
            QMessageBox.warning(self, "Σφάλμα", "Δεν βρέθηκε βάση δεδομένων")
            return

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        file_path, _ = QFileDialog.getSaveFileName(
            self,
            "Αποθήκευση Backup",
            f"payroll_backup_{timestamp}.db",
            "Database Files (*.db)"
        )

        if not file_path:
            return

        shutil.copy(DB_NAME, file_path)
        QMessageBox.information(self, "Backup", "Το backup ολοκληρώθηκε")

    def create_charts(self):
        records = self.get_filtered_records()

        if not records:
            QMessageBox.warning(self, "Σφάλμα", "Δεν υπάρχουν δεδομένα")
            return

        totals = {}

        for record in records:
            employee = record[1]
            if employee not in totals:
                totals[employee] = 0
            totals[employee] += float(record[6] or 0)

        names = list(totals.keys())
        amounts = list(totals.values())

        plt.figure(figsize=(10, 6))
        plt.bar(names, amounts)
        plt.title("Πραγματικές ώρες ανά υπάλληλο")
        plt.xlabel("Υπάλληλος")
        plt.ylabel("Ώρες")
        plt.xticks(rotation=45)
        plt.tight_layout()
        plt.show()


app = QApplication(sys.argv)

window = PayrollApp()
window.show()

sys.exit(app.exec())
